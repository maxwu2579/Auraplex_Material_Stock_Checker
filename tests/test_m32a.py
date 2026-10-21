"""Explicit cache-mode acceptance. Synthetic files and temporary databases only."""
from pathlib import Path
from contextlib import closing
from datetime import datetime
import json
import os
import shutil
import sqlite3
from unittest.mock import patch

from openpyxl import load_workbook

from checker import cached, binding, report
from checker.core import check_materials
from tests.test_m31 import UiBase, Base, digest


class CacheStorageTests(Base):
    def test_missing_database_is_not_created(self):
        path = self.base / 'absent.db'
        with self.assertRaises(cached.CacheError):
            cached.load(path)
        self.assertFalse(path.exists())

    def test_both_successful_roles_required(self):
        self.sync('stock')
        with self.assertRaises(cached.CacheError):
            cached.load(self.db)
        self.sync('bom')
        snapshot = cached.load(self.db)
        self.assertEqual([v.role for v in snapshot.versions], ['stock', 'bom'])

    def test_readonly_cache_preserves_raw_values_parity_and_database(self):
        self.write_stock([('A', None), ('B', '0+10'), ('C', 2), ('C', 3)])
        self.write_bom([('A', 2), ('B', 1), ('C', 1)])
        self.sync('stock'); self.sync('bom')
        before = digest(self.db)
        snapshot = cached.load(self.db)
        from checker.reader import ReadOnlyCache
        reader = ReadOnlyCache()
        stock = reader.read(str(self.stock), 'stock'); bom = reader.read(str(self.bom), 'bom')
        self.assertEqual(check_materials(snapshot.bom.products[0], snapshot.stock.stock_sheets[0], 3),
                         check_materials(bom.products[0], stock.stock_sheets[0], 3))
        self.assertEqual([m.quantity for m in snapshot.stock.stock_sheets[0].materials], [None, '0+10', 2, 3])
        cached.verify(snapshot)
        self.assertEqual(digest(self.db), before)

    def test_version_or_record_change_invalidates_selection(self):
        self.sync('stock'); self.sync('bom'); snapshot = cached.load(self.db)
        self.write_stock([('A', 0)]); self.sync('stock')
        with self.assertRaises(cached.CacheError): cached.verify(snapshot)
        snapshot = cached.load(self.db)
        with closing(self.inv.connect()) as conn:
            conn.execute("UPDATE stock_records SET qty_text = '999' WHERE source_version_id = ?",
                         (snapshot.versions[0].id,))
        with self.assertRaises(cached.CacheError): cached.verify(snapshot)

    def test_bad_or_foreign_schema_is_refused(self):
        for raw in (b'not sqlite',):
            path = self.base / 'bad.db'; path.write_bytes(raw)
            with self.assertRaises(cached.CacheError): cached.load(path)
        path = self.base / 'foreign.db'
        with closing(sqlite3.connect(path)) as conn:
            conn.execute('CREATE TABLE x (id)'); conn.commit()
        with self.assertRaises(cached.CacheError): cached.load(path)

    def test_incomplete_cached_rows_are_refused(self):
        self.sync('stock'); self.sync('bom')
        with closing(self.inv.connect()) as conn:
            conn.execute('DELETE FROM stock_records WHERE id = (SELECT min(id) FROM stock_records)')
        with self.assertRaises(cached.CacheReadError): cached.load(self.db)

    def test_replaced_database_expires_existing_selection_even_with_identical_ids(self):
        self.sync('stock'); self.sync('bom')
        snapshot = cached.load(self.db)
        other = self.base/'replacement.db'
        shutil.copy2(self.db, other)
        os.replace(other, self.db)
        with self.assertRaises(cached.CacheError): cached.verify(snapshot)


class CacheUiTests(UiBase):
    def setUp(self):
        super().setUp()
        self.start_app(); self.select_files(); self.check()

    def opt_in(self, answer=True):
        with patch('checker.ui.messagebox.askyesno', return_value=answer) as prompt:
            self.app.use_cached_data(); self.finish()
        return prompt

    def offline(self):
        self.stock.unlink(); self.bom.unlink()
        self.app.refresh_excel(); self.finish()
        self.assertIsNone(self.app.result)
        self.assertIsNone(self.app.cache_selection)

    def test_offline_does_not_fallback_and_declined_consent_does_not_enable(self):
        self.offline()
        self.opt_in(False)
        self.assertIsNone(self.app.cache_selection)
        self.assertIsNone(self.app.result)

    def test_offline_explicit_check_and_report_are_always_unverified(self):
        baseline = self.app.result
        self.offline()
        prompt = self.opt_in()
        self.assertIn('version', prompt.call_args.args[1])
        self.assertIn('Imported (UTC)', prompt.call_args.args[1])
        self.assertIsNone(self.app.result)
        self.check()
        self.assertEqual(self.app.result, baseline)
        self.assertFalse(self.app.result_current)
        self.assertTrue(self.app.result_exportable)
        self.assertIn(cached.LABEL, self.app.result_state_label.cget('text'))
        self.assertIn(cached.LABEL, self.app.status.get())
        self.assertIn('Using Previous Data', self.app.db_status.get())
        target = self.app.write_export(self.base / 'cached.xlsx')
        self.assertTrue(target.name.startswith('UNVERIFIED_'))
        book = load_workbook(target)
        try:
            for sheet in book:
                self.assertIn('UNVERIFIED', sheet['A1'].value)
                self.assertIn(cached.LABEL, sheet['A8'].value)
                self.assertIn('Cloud Not Verified', sheet['A8'].value)
                for row in (5, 6):
                    metadata = json.loads(sheet.cell(row, 3).comment.text)
                    self.assertIn('UNVERIFIED', metadata['scope'])
                    self.assertTrue(metadata['imported_at'])
        finally: book.close()

    def test_cached_result_cannot_enter_verified_writer(self):
        self.opt_in(); self.check()
        with self.assertRaises(binding.ExportVerificationError):
            report.save_verified_report(self.base/'no.xlsx', self.app.result, self.app.stock_book,
                                        self.app.bom_book, self.app.result_binding, self.app._export_sources())
        self.assertFalse((self.base/'no.xlsx').exists())

    def test_no_source_file_access_or_database_write_in_cached_check_export(self):
        self.opt_in(); self.check()
        before = digest(self.db)
        with patch.object(self.app.service, 'load', side_effect=AssertionError('Excel access forbidden')):
            with patch('checker.reader.fingerprint', side_effect=AssertionError('Excel fingerprint forbidden')):
                self.check()
                self.assertTrue(self.app.write_export(self.base/'cached.xlsx'))
        self.assertEqual(digest(self.db), before)

    def test_external_import_blocks_direct_and_button_exports_until_new_consent(self):
        for entry in ('direct', 'button'):
            with self.subTest(entry=entry):
                self.opt_in(); self.check()
                self.write_stock([('A', 0 if entry == 'direct' else 1)])
                self.sync('stock')
                target = self.base / (entry + '.xlsx')
                with patch('checker.ui.filedialog.asksaveasfilename', return_value=str(target)) as dialog:
                    self.app.write_export(target) if entry == 'direct' else self.app.export()
                    if entry == 'button': dialog.assert_not_called()
                self.assertFalse(self.app.result_exportable)
                self.assertIn(binding.OUTDATED, self.app.status.get())
                self.assertFalse(target.exists())
                self.assertFalse(target.with_name('UNVERIFIED_' + target.name).exists())

    def test_database_failure_during_check_is_not_old_success(self):
        self.opt_in()
        with patch('checker.cached.open_readonly', side_effect=sqlite3.OperationalError('read failed')):
            self.check()
        self.assertIsNone(self.app.result)
        self.assertIn('Status Unknown', self.app.db_status.get())

    def test_database_failure_during_export_blocks_and_invalidates(self):
        self.opt_in(); self.check()
        with patch('checker.cached.open_readonly', side_effect=sqlite3.OperationalError('read failed')):
            self.assertIsNone(self.app.write_export(self.base/'no.xlsx'))
        self.assertFalse(self.app.result_exportable)
        self.assertIn('Status Unknown', self.app.db_status.get())

    def test_change_during_consent_is_not_silently_accepted_for_check(self):
        def confirm(*args, **kwargs):
            self.write_bom([('A', 999)])
            self.sync('bom')
            return True
        with patch('checker.ui.messagebox.askyesno', side_effect=confirm):
            self.app.use_cached_data(); self.finish()
        self.check()
        self.assertIsNone(self.app.result)
        self.assertFalse(self.app.cache_authorized)
        self.assertTrue(any('changed' in e for e in self.errors))

    def test_change_during_report_generation_is_blocked(self):
        self.opt_in(); self.check()
        original = report.build_report
        def build(*args):
            book = original(*args)
            self.write_stock([('A', 0)]); self.sync('stock')
            return book
        with patch.object(report, 'build_report', side_effect=build):
            self.assertIsNone(self.app.write_export(self.base/'no.xlsx'))
        self.assertFalse(self.app.cache_authorized)
        self.assertFalse(self.app.result_exportable)
        self.assertFalse(list(self.base.glob('UNVERIFIED_*.xlsx')))
        self.assertFalse(list(self.base.glob('.mrc-export-*')))

    def test_cleanup_failure_does_not_mask_cached_version_failure(self):
        self.opt_in(); self.check()
        original = report.build_report
        def build(*args):
            book = original(*args)
            self.write_stock([('A', 0)]); self.sync('stock')
            return book
        with patch.object(report, 'build_report', side_effect=build), \
             patch.object(report.os, 'unlink', side_effect=PermissionError('locked staging')):
            self.assertIsNone(self.app.write_export(self.base/'no.xlsx'))
        self.assertFalse(self.app.result_exportable)
        self.assertIn('locked staging', self.errors[-1])
        self.assertFalse((self.base/'UNVERIFIED_no.xlsx').exists())

    def test_source_selection_change_blocks_direct_export_and_disables_button(self):
        self.opt_in(); self.check()
        self.app.stock_path.set(str(self.base/'other.xlsx'))
        self.assertIsNone(self.app.write_export(self.base/'no.xlsx'))
        self.assertTrue(self.app.export_button.instate(['disabled']))
        self.assertFalse(self.app.cache_authorized)

    def test_recheck_cannot_restore_expired_consent_without_selecting_cache(self):
        self.opt_in(); self.check()
        self.write_stock([('A', 0)]); self.sync('stock')
        self.app.write_export(self.base/'no.xlsx')
        self.check()
        self.assertFalse(self.app.result_exportable)
        self.assertFalse(self.app.cache_authorized)
        self.opt_in(); self.check()
        self.assertTrue(self.app.result_exportable)
        self.assertFalse(self.app.result_current)

    def test_cached_product_search_invalid_inputs_and_busy_serialisation(self):
        self.opt_in()
        self.app.product_text.set('product'); self.app._search()
        self.assertEqual(tuple(self.app.product_combo['values']), ('PRODUCT-1',))
        self.app.check(); self.assertIsNone(self.app.result)
        self.check('0'); self.assertIsNone(self.app.result)
        self.check()
        job = self.app.job
        self.app.busy = True
        try:
            self.app.use_cached_data(); self.app.refresh_excel(); self.app.check()
            self.assertEqual(self.app.job, job)
            self.assertIsNone(self.app.write_export(self.base/'no.xlsx'))
        finally: self.app.busy = False

    def test_cached_warning_and_controls_visible_in_small_scaled_windows(self):
        self.opt_in(); self.check()
        self.root.deiconify()
        for width, height in ((960, 600), (1280, 720), (1366, 728)):
            for scale in (1, 1.25, 1.5):
                with self.subTest(size=(width, height), scale=scale):
                    self.root.tk.call('tk', 'scaling', 96/72 * scale)
                    self.root.geometry(f'{width}x{height}')
                    self.root.update(); self.app._resize(); self.root.update()
                    self.assertGreater(self.app.tree.winfo_height(), 60)
                    for control in (self.app.mode_label, self.app.cache_button, self.app.export_button):
                        self.assertTrue(control.winfo_ismapped())
                        self.assertLessEqual(control.winfo_rooty()+control.winfo_height(),
                                             self.root.winfo_rooty()+self.root.winfo_height())
                    self.assertIn(cached.LABEL, self.app.mode_label.cget('text'))
                    self.assert_cloud_notice(self.app)

    def test_refresh_returns_to_excel_and_recalculates(self):
        self.opt_in(); self.check()
        self.write_stock([('A', 0)])
        self.app.refresh_excel(); self.finish()
        self.assertIsNone(self.app.cache_selection)
        self.assertTrue(self.app.result_current)
        self.assertEqual(self.app.result.rows[0].category, 'SHORTAGE')

    def test_failed_refresh_leaves_no_cached_fallback(self):
        self.opt_in(); self.check()
        self.stock.unlink()
        self.app.refresh_excel(); self.finish()
        self.assertIsNone(self.app.cache_selection)
        self.assertIsNone(self.app.result)

    def test_restart_never_restores_cached_consent(self):
        self.opt_in(); self.check()
        self.stock.unlink(); self.bom.unlink()
        self.start_app()
        self.assertIsNone(self.app.cache_selection)
        self.assertIsNone(self.app.result)

    def test_changed_quantity_invalidates_export(self):
        self.opt_in(); self.check()
        self.app.quantity_text.set('20')
        self.assertIsNone(self.app.result)
        self.assertIsNone(self.app.write_export(self.base/'no.xlsx'))

    def test_source_overwrite_refused_and_cached_tables_match_m2(self):
        self.opt_in(); self.check()
        before = digest(self.stock), digest(self.bom)
        self.assertIsNone(self.app.write_export(self.stock))
        self.assertEqual((digest(self.stock), digest(self.bom)), before)
        expected = report.build_report(self.app.result, self.app.stock_book, self.app.bom_book,
                                       datetime.now())
        target = self.app.write_export(self.base/'cached.xlsx')
        book = load_workbook(target)
        try:
            for sheet in book:
                self.assertEqual(list(sheet.iter_rows(min_row=15, values_only=True)),
                                 list(expected[sheet.title].iter_rows(min_row=15, values_only=True)))
        finally: book.close(); expected.close()
