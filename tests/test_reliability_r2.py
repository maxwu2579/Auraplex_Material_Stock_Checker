"""Reliability R2: background Excel export, and the --self-test failure exit.

Synthetic workbooks, temporary folders and temporary SQLite only; no real data, API or user window. The self-test
cases run app.py in separate processes with a timeout."""
from datetime import datetime
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import threading
import time
import unittest
from unittest.mock import MagicMock, patch

from openpyxl import Workbook, load_workbook

from checker import binding, cached, i18n, report as report_module
from checker.report import HEADER_ROW, build_report
from tests.fixtures import bom_rows, stock_rows, workbook
from tests.test_m31 import UiBase

ROOT = Path(__file__).resolve().parents[1]
PUBLISH = "rename" if os.name == "nt" else "link"


class Held:
    """Holds the export worker at one point until released. It runs on the worker thread and never touches Tk."""

    def __init__(self, real, after=False):
        self.real, self.after = real, after
        self.entered, self.release, self.calls = threading.Event(), threading.Event(), 0

    def __call__(self, *args, **kwargs):
        self.calls += 1
        if self.after:
            value = self.real(*args, **kwargs)
            self.entered.set(); self.release.wait(10)
            return value
        self.entered.set(); self.release.wait(10)
        return self.real(*args, **kwargs)


class ExportBase(UiBase):
    def setUp(self):
        super().setUp()
        self.start_app(); self.select_files(); self.check()
        self.target = self.base / "r2.xlsx"

    def pump(self, condition, timeout=10):
        deadline = time.monotonic() + timeout
        while time.monotonic() < deadline:
            self.root.update()
            if condition():
                return
            time.sleep(0.005)
        self.fail("condition not reached")

    def reports(self):
        # Published reports only: the source workbooks live in base/drive, staging files are checked separately.
        return sorted(p.name for p in self.base.glob("*.xlsx") if not p.name.startswith(".mrc-export-"))

    def assert_nothing_left(self):
        self.assertFalse(list(self.base.glob(".mrc-export-*")))  # no staging file that looks like a report

    def status_key(self):
        return self.app.status_message.key


class BackgroundExportTests(ExportBase):
    def test_export_runs_in_the_background_with_real_stages_and_the_same_report(self):
        held, ticks = Held(report_module.build_report), []
        with patch.object(report_module, "build_report", side_effect=held), \
                patch("checker.ui.filedialog.asksaveasfilename", return_value=str(self.target)), \
                patch.object(self.app, "_set_status", wraps=self.app._set_status) as status:
            self.app.export_button.invoke()  # returns at once
            self.assertTrue(held.entered.wait(10))
            self.root.after(0, lambda: ticks.append("tick"))
            self.pump(lambda: ticks)  # the Tk loop keeps processing events while the report is generated
            self.assertTrue(self.app.busy); self.assertTrue(self.app.export_button.instate(["disabled"]))
            self.assertEqual(self.reports(), [])
            held.release.set()
            self.finish()
        self.assertEqual([c.args[0].key for c in status.call_args_list],
                         ["st_export_checking", "st_export_building", "st_export_writing", "st_export_verifying",
                          "st_exported"])  # real stages only, no percentage
        book = load_workbook(self.target)
        expected = build_report(self.app.result, self.app.stock_book, self.app.bom_book, datetime.now())
        self.assertEqual(book.sheetnames, ["Shortage", "Exceptions"])
        for name in book.sheetnames:
            rows = lambda sheet: list(sheet.iter_rows(min_row=HEADER_ROW, values_only=True))
            self.assertEqual(rows(book[name]), rows(expected[name]))
        self.assertIn("Cloud Not Verified", book["Shortage"]["A8"].value)  # provenance row as before
        self.assertEqual(self.app.export_path.get(), str(self.target))
        self.assertTrue(self.app.export_feedback.winfo_manager())
        self.assertEqual(self.reports(), ["r2.xlsx"]); self.assert_nothing_left()
        self.assertFalse(self.errors)

    def test_repeated_clicks_and_refresh_start_no_second_task(self):
        held, dialog = Held(report_module.build_report), MagicMock(return_value=str(self.target))
        earlier = set(threading.enumerate())  # a finished test's worker may still be exiting
        with patch.object(report_module, "build_report", side_effect=held), \
                patch("checker.ui.filedialog.asksaveasfilename", dialog):
            self.app.export_button.invoke()
            self.assertTrue(held.entered.wait(10))
            job = self.app.job
            self.app.export_button.invoke(); self.app.export()
            self.assertIsNone(self.app.start_export(self.base / "second.xlsx"))
            self.assertIsNone(self.app.write_export(self.base / "third.xlsx"))
            self.app.refresh_excel()  # refused while exporting: no new job, the result stays
            self.assertEqual((dialog.call_count, self.app.job), (1, job))
            self.assertEqual(sum(t.is_alive() and t.name == "m0-export" for t in set(threading.enumerate()) - earlier), 1)
            held.release.set()
            self.finish()
        self.assertEqual(held.calls, 1)
        self.assertEqual(self.reports(), ["r2.xlsx"])
        self.assertEqual(self.status_key(), "st_exported")


class ChangeDuringExportTests(ExportBase):
    def test_an_input_change_cancels_before_publication(self):
        held = Held(report_module.build_report)
        with patch.object(report_module, "build_report", side_effect=held):
            task = self.app.start_export(self.target)
            self.assertTrue(held.entered.wait(10))
            self.app.quantity_text.set("2")  # the result is invalidated; the running export may not publish
            self.assertTrue(task.cancel.is_set())
            held.release.set()
            self.finish()
        self.assertEqual(self.reports(), []); self.assert_nothing_left()
        self.assertEqual(self.status_key(), "st_export_cancelled")
        self.assertFalse(task.accepted); self.assertFalse(self.app.export_feedback.winfo_manager())
        self.assertIsNone(self.app.result); self.assertFalse(self.errors)

    def test_a_change_after_publication_keeps_the_file_but_not_as_the_current_export(self):
        real_unlink = report_module.os.unlink
        held = Held(lambda path: self.assertTrue(self.target.exists()))  # in staging cleanup: already published

        def unlink(path):
            held(path)
            return real_unlink(path)
        with patch.object(report_module.os, "unlink", side_effect=unlink):
            task = self.app.start_export(self.target)
            self.assertTrue(held.entered.wait(10))  # published and verified for the earlier result
            self.app.quantity_text.set("2")          # the late cancellation arrives after publication
            self.assertFalse(task.cancel.cancel())   # publication came first
            held.release.set()
            self.finish()
        self.assertTrue(self.target.exists())  # never deleted
        self.assertEqual(self.app.status_message.key, "st_export_late")
        self.assertIn(str(self.target), self.app.status.get())
        self.assertFalse(task.accepted); self.assertFalse(self.app.export_feedback.winfo_manager())
        self.assertIsNone(self.app.result); self.assert_nothing_left()

    def test_write_failure_and_closing_never_report_success_or_leave_a_report(self):
        def fail(_book, path):
            Path(path).write_bytes(b"partial")
            raise OSError("synthetic disk full")
        with patch.object(Workbook, "save", fail):
            self.assertIsNone(self.app.write_export(self.target))
        self.assertIn("synthetic disk full", self.errors[-1])
        self.assertNotEqual(self.status_key(), "st_exported"); self.assertFalse(self.app.export_feedback.winfo_manager())
        self.assertEqual(self.reports(), []); self.assert_nothing_left()
        errors = len(self.errors)
        held = Held(report_module.build_report)
        with patch.object(report_module, "build_report", side_effect=held):
            task = self.app.start_export(self.target)
            self.assertTrue(held.entered.wait(10))
            worker = next(t for t in threading.enumerate() if t.name == "m0-export")
            self.app.close()
            self.assertTrue(task.cancel.is_set())
            held.release.set()
            worker.join(10)
        self.assertFalse(worker.is_alive())
        self.assertEqual(self.reports(), []); self.assert_nothing_left()
        self.assertFalse(task.accepted); self.assertEqual(len(self.errors), errors)


class HoldAfterFinalCheck:
    """Holds the worker right after its n-th call of `real` (the final verification), before the publication
    boundary; calls on the Tk thread are not counted. Runs on the worker and never touches Tk."""

    def __init__(self, real, nth=2):
        self.real, self.nth, self.count = real, nth, 0
        self.entered, self.release = threading.Event(), threading.Event()

    def __call__(self, *args, **kwargs):
        value = self.real(*args, **kwargs)
        if threading.current_thread() is not threading.main_thread():
            self.count += 1
            if self.count == self.nth:
                self.entered.set(); self.release.wait(10)
        return value


class PublicationBoundaryTests(ExportBase):
    def test_cancellation_after_the_final_check_but_before_publication_wins(self):
        held = HoldAfterFinalCheck(binding.verify_export)  # 1: pre-check, 2: final check right before publishing
        with patch.object(binding, "verify_export", side_effect=held):
            task = self.app.start_export(self.target)
            self.assertTrue(held.entered.wait(10))
            self.assertFalse(self.target.exists())
            began = time.monotonic()
            self.app.quantity_text.set("2")  # Tk thread: invalidated, cancellation takes the boundary first
            self.assertLess(time.monotonic() - began, 1)  # the boundary is free: no wait on the Tk thread
            self.assertTrue(task.cancel.is_set())
            held.release.set()
            self.finish()
        self.assertEqual(self.reports(), []); self.assert_nothing_left()
        self.assertIsNone(task.cancel.published); self.assertIsNone(task.written)
        self.assertEqual(self.status_key(), "st_export_cancelled"); self.assertFalse(self.errors)

    def test_a_publication_under_way_wins_and_the_cancellation_waits_for_it(self):
        real, entered, release = getattr(report_module.os, PUBLISH), threading.Event(), threading.Event()

        def publish(*args):  # inside the boundary, before the actual rename/link
            entered.set(); release.wait(10)
            return real(*args)
        outcome = []
        with patch.object(report_module.os, PUBLISH, side_effect=publish):
            task = self.app.start_export(self.target)
            self.assertTrue(entered.wait(10))
            canceller = threading.Thread(target=lambda: outcome.append(task.cancel.cancel()))
            canceller.start()
            canceller.join(0.3)
            self.assertTrue(canceller.is_alive())  # serialised: it cannot slip in before the rename
            release.set()
            canceller.join(10)
            self.app.quantity_text.set("2")
            self.finish()
        self.assertEqual(outcome, [False])  # publication came first
        self.assertTrue(self.target.exists())
        self.assertEqual(task.cancel.published, self.target)
        self.assertEqual(self.status_key(), "st_export_late"); self.assertFalse(task.accepted)
        self.assert_nothing_left()

    def test_a_cancelled_export_reports_a_staging_file_it_could_not_remove(self):
        real_unlink, held = report_module.os.unlink, Held(report_module.build_report)

        def unlink(path):
            if Path(path).name.startswith(".mrc-export-"):
                raise PermissionError(13, "synthetic staging file locked")
            return real_unlink(path)
        with patch.object(report_module, "build_report", side_effect=held), \
                patch.object(report_module.os, "unlink", side_effect=unlink):
            task = self.app.start_export(self.target)
            self.assertTrue(held.entered.wait(10))
            self.app.quantity_text.set("2")
            held.release.set()
            self.finish()
        leftovers = list(self.base.glob(".mrc-export-*"))
        self.assertEqual(len(leftovers), 1)
        try:
            message = self.errors[-1]
            for fact in (str(leftovers[0]), "synthetic staging file locked", "cancelled before the report was published",
                         "current result was not changed"):
                self.assertIn(fact, message)
            self.assertIn(str(leftovers[0]), self.app.status.get())
            self.assertIn(i18n.tr("st_export_cancelled"), self.app.status.get())
            self.assertEqual(self.reports(), []); self.assertFalse(task.accepted)
            self.assertIsNone(self.app.result); self.assertFalse(self.app.export_feedback.winfo_manager())
        finally:
            for path in leftovers:
                path.unlink()  # only this test's own synthetic staging file


class CachedExportTests(ExportBase):
    def setUp(self):
        super().setUp()
        with patch("checker.ui.messagebox.askyesno", return_value=True):
            self.app.use_cached_data(); self.finish()
        self.check()
        self.assertIsInstance(self.app.result_binding, cached.CacheBinding)

    def test_cached_export_stays_unverified_and_a_revoked_authorisation_cancels_it(self):
        written = self.app.write_export(self.target)
        self.assertEqual(Path(written).name, "UNVERIFIED_r2.xlsx")
        book = load_workbook(written)
        for sheet in book:
            self.assertIn("UNVERIFIED", sheet["A1"].value)
            self.assertIn(cached.LABEL, sheet["A8"].value)
        self.assertEqual(self.status_key(), "st_cache_exported")
        held = Held(report_module.build_report)
        with patch.object(report_module, "build_report", side_effect=held):
            task = self.app.start_export(self.base / "revoked.xlsx")
            self.assertTrue(held.entered.wait(10))
            self.app._leave_cache()  # the explicit cache authorisation ends while the export runs
            self.assertTrue(task.cancel.is_set())
            held.release.set()
            self.finish()
        self.assertEqual(self.reports(), ["UNVERIFIED_r2.xlsx"]); self.assert_nothing_left()
        self.assertEqual(self.status_key(), "st_export_cancelled"); self.assertFalse(task.accepted)

    def test_revoking_cache_authorisation_after_the_final_check_prevents_publication(self):
        held = HoldAfterFinalCheck(cached.verify)  # 1: before building, 2: final check right before publishing
        with patch.object(cached, "verify", side_effect=held):
            task = self.app.start_export(self.target)
            self.assertTrue(held.entered.wait(10))
            self.app._leave_cache()  # Tk thread: authorisation revoked between the final check and publication
            self.assertTrue(task.cancel.is_set())
            held.release.set()
            self.finish()
        self.assertEqual(self.reports(), []); self.assert_nothing_left()
        self.assertIsNone(task.cancel.published)
        self.assertEqual(self.status_key(), "st_export_cancelled")


class SelfTestExitTests(unittest.TestCase):
    """app.py --self-test in its own process: a hang would fail the timeout, a failure must exit non-zero."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix="_MRC_R2_SELFTEST_")
        self.base = Path(self.tmp.name)
        self.stock, self.bom = self.base / "in" / "storage.xlsx", self.base / "in" / "bom.xlsx"
        self.stock.parent.mkdir()
        workbook(self.stock, [("Stock", stock_rows([("A", 100), ("B", 0)]))])
        workbook(self.bom, [("PRODUCT-1", bom_rows([("A", 2), ("B", 1)]))])

    def tearDown(self):
        self.tmp.cleanup()

    def run_self_test(self, output):
        env = dict(os.environ, LOCALAPPDATA=str(self.base / "localappdata"), APPDATA=str(self.base / "appdata"))
        began = time.monotonic()
        done = subprocess.run([sys.executable, str(ROOT / "app.py"), "--self-test", str(self.stock), str(self.bom),
                               str(output)], cwd=ROOT, env=env, capture_output=True, text=True, timeout=120)
        return done, time.monotonic() - began

    def test_a_missing_output_folder_is_created_and_success_exits_0(self):
        output = self.base / "new" / "nested"
        done, _ = self.run_self_test(output)
        self.assertEqual(done.returncode, 0, done.stderr)
        result = json.loads((output / "self_test.json").read_text(encoding="utf-8"))
        self.assertTrue(result["passed"], result["errors"])
        self.assertTrue(Path(result["export"]).is_file())
        self.assertFalse((self.base / "localappdata").exists())  # only the given folder was used

    def test_an_unusable_folder_or_result_file_exits_non_zero(self):
        not_a_folder = self.base / "occupied"
        not_a_folder.write_bytes(b"unrelated")
        done, seconds = self.run_self_test(not_a_folder)
        self.assertEqual(done.returncode, 2)
        self.assertIn("OUTPUT_DIR is unusable", done.stderr)
        self.assertEqual(not_a_folder.read_bytes(), b"unrelated")
        self.assertLess(seconds, 60)
        output = self.base / "out"
        (output / "self_test.json").mkdir(parents=True)  # the result file itself cannot be written
        done, _ = self.run_self_test(output)
        self.assertEqual(done.returncode, 2)  # the window closed and the process ended within the timeout
        self.assertIn("could not be written", done.stderr)
        self.assertTrue((output / "self_test.json").is_dir())


if __name__ == "__main__":
    unittest.main()
