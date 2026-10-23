"""P1 export gate regressions. Only synthetic workbooks and private temporary output are used."""
from dataclasses import replace
from datetime import datetime
import json
import os
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook

from checker import binding, report as report_module
from checker.core import DataReadError
from checker.report import HEADER_ROW, build_report, save_verified_report
from tests.test_m31 import UiBase, digest
from tests.fixtures import stock_rows, workbook


class ExportFreshnessTests(UiBase):
    def setUp(self):
        super().setUp()
        self.app = self.start_app()
        self.select_files()
        self.check()
        self.target = self.base / 'export.xlsx'

    def assert_refused(self):
        self.assertFalse(self.app.result_current)
        self.assertTrue(self.app.export_button.instate(['disabled']))
        self.assertIn(binding.OUTDATED, self.app.result_state_label.cget('text'))
        self.assertIn(binding.OUTDATED, self.app.status.get())
        self.assertIn('Cloud Not Verified', self.app.status.get())
        self.assertFalse(self.target.exists())
        self.assertFalse(list(self.base.glob('.mrc-export-*')))
        self.assertTrue(self.errors)

    def test_external_storage_save_after_check_blocks_direct_export(self):
        self.assertEqual(self.app.result.rows[0].category, 'ENOUGH')
        self.write_stock([('A', 0)])
        self.assertIsNone(self.app.write_export(self.target))
        self.assert_refused()

    def test_external_bom_save_after_check_blocks_direct_export(self):
        self.write_bom([('A', 200)])
        self.assertIsNone(self.app.write_export(self.target))
        self.assert_refused()

    def test_ui_preflight_refuses_before_opening_dialog(self):
        self.write_stock([('A', 0)])
        with patch('checker.ui.filedialog.asksaveasfilename', side_effect=AssertionError('must not open dialog')):
            self.app.export_button.invoke()
        self.assert_refused()

    def test_change_while_save_dialog_is_open_is_caught_at_write_boundary(self):
        def choose(**_):
            self.write_bom([('A', 200)])
            return str(self.target)
        with patch('checker.ui.filedialog.asksaveasfilename', side_effect=choose):
            self.app.export()
        self.assert_refused()

    def test_changed_result_during_dialog_is_not_exported_under_old_name(self):
        def choose(**_):
            self.app.quantity_text.set('2')
            self.app.check(); self.finish()
            return str(self.target)
        with patch('checker.ui.filedialog.asksaveasfilename', side_effect=choose):
            self.app.export()
        self.assertFalse(self.target.exists())
        self.assertIn('result changed while you chose the export location', self.errors[-1])  # B1: English interface
        self.assertEqual(self.app.result.production_count, 2)

    def test_selected_source_path_must_match_check_even_for_identical_bytes(self):
        other = self.base / 'other.xlsx'
        other.write_bytes(self.stock.read_bytes())
        self.app.stock_path.set(str(other))
        self.assertIsNone(self.app.write_export(self.target))
        self.assert_refused()

    def test_unavailable_source_is_not_treated_as_historical_success(self):
        self.stock.unlink()
        self.assertIsNone(self.app.write_export(self.target))
        self.assert_refused()
        self.assertIn('unavailable', self.app.status.get())

    def test_file_permission_failure_does_not_fall_back_to_database_or_cache(self):
        with patch.object(binding, 'fingerprint', side_effect=DataReadError('Permission denied')):
            self.assertIsNone(self.app.write_export(self.target))
        self.assert_refused()
        self.assertIn('Permission denied', self.app.status.get())

    def test_atomic_source_replacement_with_invalid_workbook_is_refused(self):
        replacement = self.base / 'replacement.xlsx'
        replacement.write_bytes(b'not a ZIP workbook')
        os.replace(replacement, self.bom)
        self.assertIsNone(self.app.write_export(self.target))
        self.assert_refused()

    def test_changed_content_is_detected_even_when_metadata_is_preserved(self):
        before = self.app.result_binding.versions['stock']
        self.write_stock([('A', 101), ('B', 5), ('C', 0)])
        changed = binding.fingerprint(self.stock)
        real_fingerprint = binding.fingerprint
        def same_metadata(path):
            return before[:6] + (changed[6],) if Path(path) == self.stock else real_fingerprint(path)
        with patch.object(binding, 'fingerprint', side_effect=same_metadata):
            self.assertIsNone(self.app.write_export(self.target))
        self.assert_refused()

    def test_storage_changing_while_bom_is_fingerprinted_is_refused(self):
        real_fingerprint = binding.fingerprint
        calls = []
        def observe(path):
            value = real_fingerprint(path)
            calls.append(str(path))
            if Path(path) == self.bom and len(calls) == 2:
                self.write_stock([('A', 0)])
            return value
        with patch.object(binding, 'fingerprint', side_effect=observe):
            self.assertIsNone(self.app.write_export(self.target))
        self.assertEqual(len(calls), 4)  # fixed two passes, no unbounded retries
        self.assert_refused()

    def test_later_matching_observation_does_not_hide_an_earlier_failure(self):
        real_fingerprint = binding.fingerprint
        calls = []
        def observe(path):
            calls.append(path)
            if len(calls) == 1:
                raise DataReadError('temporarily unavailable')
            return real_fingerprint(path)
        with patch.object(binding, 'fingerprint', side_effect=observe):
            self.assertIsNone(self.app.write_export(self.target))
        self.assertEqual(len(calls), 4)
        self.assert_refused()

    def test_change_during_report_generation_never_publishes_a_final_report(self):
        real_build = report_module.build_report
        def build(*args):
            book = real_build(*args)
            self.write_stock([('A', 0)])
            return book
        with patch.object(report_module, 'build_report', side_effect=build):
            self.assertIsNone(self.app.write_export(self.target))
        self.assert_refused()

    def test_change_after_staging_save_never_publishes_a_final_report(self):
        real_save = Workbook.save
        def save(book, path):
            real_save(book, path)
            self.write_bom([('A', 200)])
        with patch.object(Workbook, 'save', save):
            self.assertIsNone(self.app.write_export(self.target))
        self.assert_refused()

    def test_partial_write_failure_cleans_only_staging_file(self):
        self.target.write_bytes(b'existing report')
        def fail(_book, path):
            Path(path).write_bytes(b'partial')
            raise OSError('disk full')
        with patch.object(Workbook, 'save', fail):
            self.assertIsNone(self.app.write_export(self.target))
        self.assertEqual(self.target.read_bytes(), b'existing report')
        self.assertFalse((self.base / 'export (2).xlsx').exists())
        self.assertFalse(list(self.base.glob('.mrc-export-*')))
        self.assertIn('disk full', self.errors[-1])

    def test_cleanup_failure_cannot_hide_changed_source_from_ui(self):
        real_build = report_module.build_report
        def build(*args):
            book = real_build(*args)
            self.write_stock([('A', 0)])
            return book
        with patch.object(report_module, 'build_report', side_effect=build):
            with patch.object(report_module.os, 'unlink', side_effect=PermissionError('staging file locked')):
                self.assertIsNone(self.app.write_export(self.target))
        leftovers = list(self.base.glob('.mrc-export-*'))
        self.assertEqual(len(leftovers), 1)
        for path in leftovers:
            path.unlink()  # release the injected failure; clean only this test's staging file
        self.assert_refused()
        self.assertIn('staging file locked', self.errors[-1])
        self.assertIn(str(leftovers[0]), self.errors[-1])

    def test_cleanup_failure_preserves_verification_exception_at_writer_boundary(self):
        real_build = report_module.build_report
        def build(*args):
            book = real_build(*args)
            self.write_bom([('A', 200)])
            return book
        with patch.object(report_module, 'build_report', side_effect=build):
            with patch.object(report_module.os, 'unlink', side_effect=PermissionError('staging file locked')):
                with self.assertRaises(binding.ExportVerificationError) as caught:
                    save_verified_report(self.target, self.app.result, self.app.stock_book, self.app.bom_book,
                                         self.app.result_binding, self.app._export_sources())
        self.assertEqual(caught.exception.binding.changed, ('bom',))
        self.assertIn('staging file locked', caught.exception.cleanup_warning)
        self.assertFalse(self.target.exists())

    def test_cleanup_failure_preserves_original_write_error_and_reports_leftover(self):
        def fail(_book, path):
            Path(path).write_bytes(b'partial')
            raise OSError('disk full')
        with patch.object(Workbook, 'save', fail):
            with patch.object(report_module.os, 'unlink', side_effect=PermissionError('staging file locked')):
                self.assertIsNone(self.app.write_export(self.target))
        self.assertFalse(self.target.exists())
        self.assertIn('disk full', self.errors[-1])
        self.assertIn('staging file locked', self.errors[-1])

    def test_publish_collision_never_overwrites_another_report(self):
        real_publish = os.rename if os.name == 'nt' else os.link
        calls = []
        def publish(source, target):
            calls.append(target)
            if len(calls) == 1:
                Path(target).write_bytes(b'created by other application')
            return real_publish(source, target)
        name = 'rename' if os.name == 'nt' else 'link'
        with patch.object(report_module.os, name, side_effect=publish):
            written = self.app.write_export(self.target)
        self.assertEqual(self.target.read_bytes(), b'created by other application')
        self.assertEqual(written, self.base / 'export (2).xlsx')
        self.assertFalse(list(self.base.glob('.mrc-export-*')))

    def test_publish_collision_retries_are_bounded(self):
        name = 'rename' if os.name == 'nt' else 'link'
        with patch.object(report_module.os, name, side_effect=FileExistsError('collision')) as publish:
            self.assertIsNone(self.app.write_export(self.target))
        self.assertEqual(publish.call_count, 10)
        self.assertFalse(self.target.exists())
        self.assertFalse(list(self.base.glob('.mrc-export-*')))

    def test_lower_export_boundary_rejects_a_different_workbook_snapshot(self):
        old = self.app.stock_book
        mismatched = replace(old, fingerprint=old.fingerprint[:6] + ((),))
        with self.assertRaises(binding.ExportVerificationError):
            save_verified_report(self.target, self.app.result, mismatched, self.app.bom_book,
                                 self.app.result_binding, self.app._export_sources())
        self.assertFalse(self.target.exists())

    def test_failed_binding_cannot_be_restored_without_check(self):
        self.write_stock([('A', 0)])
        self.assertIsNone(self.app.write_export(self.target))
        with patch.object(binding, 'fingerprint', side_effect=AssertionError('must require CHECK first')):
            self.assertIsNone(self.app.write_export(self.target))
        self.assert_refused()
        self.check()
        self.assertTrue(self.app.write_export(self.target))
        self.assertEqual(self.app.result.rows[0].category, 'SHORTAGE')

    def test_success_keeps_m2_tables_sources_database_and_records_provenance(self):
        before = digest(self.stock), digest(self.bom), digest(self.db)
        parses = self.app.service.cache.parse_count
        baseline = build_report(self.app.result, self.app.stock_book, self.app.bom_book, datetime.now())
        try:
            written = self.app.write_export(self.target)
            self.assertTrue(written)
            with open(written, 'rb') as stream:
                book = load_workbook(stream)
                try:
                    self.assertEqual(book.sheetnames, ['Shortage', 'Exceptions'])
                    for sheet in book:
                        expected = list(baseline[sheet.title].iter_rows(min_row=HEADER_ROW, values_only=True))
                        self.assertEqual(list(sheet.iter_rows(min_row=HEADER_ROW, values_only=True)), expected)
                        self.assertIn('Cloud Not Verified', sheet['A8'].value)
                        for role, row in (('stock', 5), ('bom', 6)):
                            metadata = json.loads(sheet.cell(row, 3).comment.text)
                            self.assertEqual(metadata['path'], self.app.result_binding.versions[role][0])
                            self.assertEqual(len(metadata['version_id']), 64)
                            self.assertGreaterEqual(metadata['export_verified_at'], metadata['check_verified_at'])
                finally:
                    book.close()
        finally:
            baseline.close()
        self.assertEqual((digest(self.stock), digest(self.bom), digest(self.db)), before)
        self.assertEqual(self.app.service.cache.parse_count, parses)

    def test_export_does_not_read_sqlite_when_local_excel_is_verified(self):
        with patch.object(self.app.inventory, 'freshness', side_effect=RuntimeError('database unavailable')):
            self.assertTrue(self.app.write_export(self.target))
        self.assert_cloud_notice(self.app)

    def test_busy_check_cannot_be_bypassed_by_direct_export(self):
        self.app.busy = True
        try:
            self.assertIsNone(self.app.write_export(self.target))
            with patch('checker.ui.filedialog.asksaveasfilename', side_effect=AssertionError('busy')):
                self.app.export()
        finally:
            self.app.busy = False
        self.assertFalse(self.target.exists())

    def test_large_media_export_does_not_reparse_sources_or_read_images(self):
        workbook(self.stock, [('Stock', stock_rows([('A', 100)]))], media=b'0' * (101 << 20))
        self.check()
        self.assertGreater(self.stock.stat().st_size, 100 << 20)
        real_read = ZipFile.read
        def read(archive, name, *args, **kwargs):
            self.assertFalse(str(name).startswith('xl/media/'), 'image content must not be loaded')
            return real_read(archive, name, *args, **kwargs)
        with patch.object(self.app.service.cache, 'read', side_effect=AssertionError('no source reparse')):
            with patch.object(ZipFile, 'read', read):
                self.assertTrue(self.app.write_export(self.target))
