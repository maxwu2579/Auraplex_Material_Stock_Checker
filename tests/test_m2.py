"""M2 regression tests: approved stock rules, categories, sorting, Product ID selection and the XLSX report.

Synthetic files only. DriveMountTests run the file-behaviour checks on a real mounted folder when the
environment variable MRC_TEST_MOUNT_DIR names one (for example the Google Drive for Desktop root); they
create and delete their own temporary sub-folder there and are skipped otherwise.
"""
import ctypes
from datetime import datetime
import os
from pathlib import Path
import tempfile
import time
import tkinter as tk
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from checker import reader
from checker.core import DataReadError, check_materials, search_products, select_product
from checker.reader import ReadOnlyCache
from checker.report import ALL_HEADER, EXCEPTION_HEADER, HEADER_ROW, ReportError, default_filename, safe_target, save_report
from checker.report_style import METADATA_START, METADATA_VALUE_COLUMN, NOTE_ROW, RED, YELLOW
from checker.service import CheckerService
from checker.ui import CheckerApp
from tests.fixtures import workbook, stock_rows, bom_rows

STOCK = [('SHORT', 12), ('ENOUGH', 100), ('EXACT', 20), ('BLANK', None), ('TEXT', '0+10'), ('ZERO', 0),
         ('DUP', 3), ('DUP', 7), ('TWICE', 50), ('BADQTY', 5), ('=FORMULA', 9)]
BOM = [('ENOUGH', 1), ('TWICE', 1), ('SHORT', 2), ('MISSING', 1), ('BLANK', 1), ('TEXT', 1), ('ZERO', 1), ('DUP', 1),
       ('TWICE', 2), ('BADQTY', 'x'), ('EXACT', 2), ('=FORMULA', 1), ('ALSO MISSING', 0)]


class Files(unittest.TestCase):
    folder = None  # None: system temp directory

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(dir=self.folder, prefix='_MRC_M2_TEST_')
        self.base = Path(self.tmp.name)
        self.stock, self.bom = self.base / 'storage.xlsx', self.base / 'bom.xlsx'
        workbook(self.stock, [('Stock', stock_rows(STOCK))])
        workbook(self.bom, [('DEMO-200-A-40W ', bom_rows(BOM)), ('DEMO-200-A-25W', bom_rows([('ENOUGH', 1)])),
                            ('KIT-115', bom_rows([('SHORT', 1)])), ('BOM TEMPLATE', bom_rows([]))])
        self.cache = ReadOnlyCache()

    def tearDown(self):
        self.tmp.cleanup()

    def books(self):
        return self.cache.read(str(self.stock), 'stock'), self.cache.read(str(self.bom), 'bom')

    def result(self, count=10):
        stock, bom = self.books()
        return check_materials(bom.products[0], stock.stock_sheets[0], count)

    def row(self, code, count=10):
        return next(r for r in self.result(count).rows if r.bom.code == code)

    def one(self, stock_value, bom_qty=2, count=10):
        """A single material 'A' with the given raw stock value and BOM quantity."""
        workbook(self.stock, [('Stock', stock_rows([('A', stock_value)]))])
        workbook(self.bom, [('P', bom_rows([('A', bom_qty)]))])
        return self.result(count).rows[0]


class RuleTests(Files):
    def test_01_valid_shortage(self):
        r = self.row('SHORT')
        self.assertEqual((r.required, r.stock_used, r.shortage), (20, 12, 8))
        self.assertEqual((r.category, r.data_status, r.stock_assumed), ('SHORTAGE', 'NORMAL', False))

    def test_02_valid_enough_including_required_equal_to_stock(self):
        for code, required, stock in (('ENOUGH', 10, 100), ('EXACT', 20, 20)):
            r = self.row(code)
            self.assertEqual((r.required, r.stock_used, r.shortage, r.category, r.data_status), (required, stock, 0, 'ENOUGH', 'NORMAL'))

    def test_03_material_not_found_uses_zero_and_is_exception(self):
        r = self.row('MISSING')
        self.assertEqual((r.required, r.stock_used, r.stock_assumed, r.shortage), (10, 0, True, 10))
        self.assertEqual((r.category, r.data_status, r.stock), ('EXCEPTION', 'STOCK_NOT_FOUND / ASSUMED 0', ()))

    def test_04_blank_stock_uses_zero_and_is_exception(self):
        r = self.row('BLANK')
        self.assertEqual((r.stock_used, r.stock_assumed, r.shortage), (0, True, 10))
        self.assertEqual((r.category, r.data_status, r.original_stock), ('EXCEPTION', 'BLANK_STOCK / ASSUMED 0', '(blank)'))

    def test_05_text_and_other_invalid_stock_use_zero_keep_raw_value_and_are_not_evaluated(self):
        r = self.row('TEXT')
        self.assertEqual((r.stock_used, r.stock_assumed, r.shortage, r.original_stock), (0, True, 10, '0+10'))
        self.assertEqual((r.category, r.data_status), ('EXCEPTION', 'INVALID_STOCK / ASSUMED 0'))
        for raw in ('N/A', 'unknown', '10', -1, ('f', '0+10'), ('e', '#VALUE!'), ('b', '1')):
            with self.subTest(raw=raw):
                r = self.one(raw)
                self.assertEqual((r.stock_used, r.stock_assumed, r.shortage, r.category), (0, True, 20, 'EXCEPTION'))
                self.assertEqual(r.data_status, 'INVALID_STOCK / ASSUMED 0'); self.assertTrue(r.original_stock)

    def test_06_numeric_zero_is_a_normal_shortage_not_an_exception(self):
        r = self.row('ZERO')
        self.assertEqual((r.stock_used, r.stock_assumed, r.shortage), (0, False, 10))
        self.assertEqual((r.category, r.data_status), ('SHORTAGE', 'NORMAL'))

    def test_07_duplicate_stock_is_exception_without_sum_or_first_record(self):
        r = self.row('DUP')
        self.assertEqual((r.category, r.data_status), ('EXCEPTION', 'DUPLICATE_STOCK'))
        self.assertEqual((r.stock_used, r.stock_assumed, r.shortage, r.delta), (None, False, None, None))
        self.assertEqual(([s.quantity for s in r.stock], r.original_stock), ([3, 7], '3 | 7')); self.assertEqual(r.required, 10)

    def test_08_duplicate_bom_rows_are_kept_separately_and_not_merged(self):
        rows = [r for r in self.result().rows if r.bom.code == 'TWICE']
        self.assertEqual([r.required for r in rows], [10, 20]); self.assertEqual([r.bom.source.row for r in rows], [23, 30])
        self.assertTrue(all((r.category, r.data_status, r.shortage) == ('EXCEPTION', 'DUPLICATE_BOM', None) for r in rows))
        self.assertTrue(all(r.stock_used == 50 and not r.stock_assumed for r in rows))  # the one stock value is still shown

    def test_09_invalid_bom_quantity(self):
        r = self.row('BADQTY')
        self.assertEqual((r.required, r.shortage, r.category, r.data_status, r.bom.quantity), (None, None, 'EXCEPTION', 'INVALID_BOM_QTY', 'x'))
        for raw in (None, 'two', ('f', '1+1'), ('e', '#REF!'), ('b', '1'), -2, 0):
            with self.subTest(raw=raw):
                r = self.one(5, bom_qty=raw)
                self.assertEqual((r.required, r.shortage, r.category, r.data_status), (None, None, 'EXCEPTION', 'INVALID_BOM_QTY'))

    def test_09b_several_exceptions_on_one_row_are_all_reported(self):
        r = self.row('ALSO MISSING')  # zero BOM QTY and not in Storage
        self.assertEqual((r.category, r.exceptions, r.shortage), ('EXCEPTION', ('INVALID_BOM_QTY', 'STOCK_NOT_FOUND'), None))
        self.assertEqual(r.data_status, 'INVALID_BOM_QTY / STOCK_NOT_FOUND / ASSUMED 0')

    def test_10_every_row_has_one_category_and_rows_sort_red_yellow_white(self):
        result = self.result()
        self.assertEqual(len(result.rows), len(BOM)); self.assertEqual([r.bom.code for r in result.rows], [c for c, _ in BOM])  # BOM order kept
        self.assertTrue(all(r.category in ('SHORTAGE', 'EXCEPTION', 'ENOUGH') for r in result.rows))
        self.assertTrue(all((r.category == 'EXCEPTION') == bool(r.exceptions) for r in result.rows))
        ordered = [(r.category, r.bom.code) for r in result.sorted_rows]
        # '=FORMULA' is an ordinary text code that merely starts with "="; it is a normal row (stock 9, required 10).
        self.assertEqual(ordered, [('SHORTAGE', '=FORMULA'), ('SHORTAGE', 'SHORT'), ('SHORTAGE', 'ZERO'),
                                   ('EXCEPTION', 'ALSO MISSING'), ('EXCEPTION', 'BADQTY'), ('EXCEPTION', 'BLANK'),
                                   ('EXCEPTION', 'DUP'), ('EXCEPTION', 'MISSING'), ('EXCEPTION', 'TEXT'), ('EXCEPTION', 'TWICE'), ('EXCEPTION', 'TWICE'),
                                   ('ENOUGH', 'ENOUGH'), ('ENOUGH', 'EXACT')])
        self.assertEqual([result.category_count(c) for c in ('SHORTAGE', 'EXCEPTION', 'ENOUGH')], [3, 8, 2])
        self.assertEqual(result.calculated_shortage_rows, 6)  # =FORMULA, SHORT, ZERO + MISSING, BLANK, TEXT

    def test_10b_exception_with_calculated_shortage_never_moves_into_shortage(self):
        r = self.row('MISSING', count=15)
        self.assertEqual((r.shortage, r.category), (15, 'EXCEPTION'))

    def test_12_production_quantity_must_be_a_positive_integer(self):
        stock, bom = self.books()
        for bad in (0, -1, 1.5, True, '3'):
            with self.subTest(bad=bad), self.assertRaises(ValueError):
                check_materials(bom.products[0], stock.stock_sheets[0], bad)
        self.assertEqual(self.row('SHORT', count=1).required, 2)


class ProductIdTests(Files):
    def test_11_product_ids_come_from_bom_sheets_and_keep_the_source_mapping(self):
        products = self.books()[1].products
        self.assertEqual([p.display_name for p in products], ['DEMO-200-A-40W', 'DEMO-200-A-25W', 'KIT-115'])  # template excluded
        first = products[0]
        self.assertEqual((first.sheet_name, first.file_path, first.header_row), ('DEMO-200-A-40W ', str(self.bom.resolve()), 21))
        self.assertEqual(first.materials[0].source.sheet, 'DEMO-200-A-40W ')  # original Sheet name, trailing space kept

    def test_11_partial_search_is_case_insensitive_but_only_a_full_id_selects(self):
        products = self.books()[1].products
        self.assertEqual([p.display_name for p in search_products(products, 'demo-200')], ['DEMO-200-A-40W', 'DEMO-200-A-25W'])
        self.assertEqual([p.display_name for p in search_products(products, '25w')], ['DEMO-200-A-25W'])
        for text in ('demo-200', '25w', 'DEMO', 'NOT-A-PRODUCT', '', '   '):
            with self.subTest(text=text):
                self.assertIsNone(select_product(products, text))
        self.assertEqual(select_product(products, 'demo-200-a-40w').sheet_name, 'DEMO-200-A-40W ')
        self.assertEqual(select_product(products, 'DEMO-200-A-40W').sheet_name, 'DEMO-200-A-40W ')

    def test_11_lookalike_sheet_names_stay_distinct_product_ids(self):
        workbook(self.bom, [('Same', bom_rows([('A', 1)])), ('Same ', bom_rows([('B', 1)]))])
        products = self.books()[1].products
        self.assertEqual(len({p.display_name for p in products}), 2); self.assertIsNone(select_product(products, 'Same'))
        self.assertEqual(select_product(products, products[1].display_name).sheet_name, 'Same ')


class ReportTests(Files):
    def setUp(self):
        super().setUp()
        self.stock_book, self.bom_book = self.books()
        self.checked = check_materials(self.bom_book.products[0], self.stock_book.stock_sheets[0], 10)
        self.before = (self.stock.read_bytes(), self.bom.read_bytes())

    def save(self, name='report.xlsx'):
        return save_report(self.base / name, self.checked, self.stock_book, self.bom_book, datetime(2026, 10, 7, 9, 30, 5))

    def table(self, sheet):
        return [dict(zip([c.value for c in sheet[HEADER_ROW]], row)) for row in sheet.iter_rows(min_row=HEADER_ROW + 1, values_only=True)]

    def test_13_export_creates_a_new_xlsx_with_exactly_two_sheets(self):
        target = self.save(); book = load_workbook(target)
        self.assertEqual(target, self.base / 'report.xlsx'); self.assertEqual(book.sheetnames, ['Shortage', 'Exceptions'])
        for sheet in book:
            top = {sheet.cell(row=n, column=1).value: sheet.cell(row=n, column=METADATA_VALUE_COLUMN).value for n in range(METADATA_START, METADATA_START + 5)}
            self.assertEqual(top, {'Product ID': 'DEMO-200-A-40W', 'Production Quantity': 10, 'Storage File': str(self.stock.resolve()),
                                   'BOM File': str(self.bom.resolve()), 'Generated Time': '2026-10-07 09:30:05'})
            self.assertEqual(sheet.freeze_panes, f'A{HEADER_ROW + 1}'); self.assertTrue(sheet.auto_filter.ref.startswith(f'A{HEADER_ROW}:'))
            self.assertNotIn('READY', str(sheet.cell(row=NOTE_ROW, column=1).value).upper())
        self.assertEqual([c.value for c in book['Shortage'][HEADER_ROW]], list(ALL_HEADER))
        self.assertEqual([c.value for c in book['Exceptions'][HEADER_ROW]], list(EXCEPTION_HEADER))

    def test_14_shortage_sheet_contains_all_rows_with_values(self):
        rows = self.table(load_workbook(self.save())['Shortage'])
        self.assertEqual(len(rows), len(BOM)); self.assertEqual([r['No.'] for r in rows], list(range(1, len(BOM) + 1)))
        by = {r['Material ID']: r for r in rows if r['Material ID'] != 'TWICE'}
        short, missing, dup, bad = by['SHORT'], by['MISSING'], by['DUP'], by['BADQTY']
        self.assertEqual((short['BOM QTY'], short['Production Quantity'], short['Required'], short['Stock'], short['Shortage'], short['Result'],
                          short['Data Status'], short['Source Sheet'], short['Source Row']), (2, 10, 20, 12, 8, 'SHORTAGE', 'NORMAL', 'DEMO-200-A-40W ', 24))
        self.assertEqual((missing['Stock'], missing['Shortage'], missing['Result'], missing['Data Status']), (0, 10, 'EXCEPTION', 'STOCK_NOT_FOUND / ASSUMED 0'))
        self.assertEqual((dup['Stock'], dup['Shortage'], dup['Original Stock Value'], dup['Data Status']), (None, None, '3 | 7', 'DUPLICATE_STOCK'))
        self.assertEqual((bad['BOM QTY'], bad['Required'], bad['Shortage']), ('x', None, None))
        self.assertEqual((by['TEXT']['Original Stock Value'], by['TEXT']['Stock']), ('0+10', 0)); self.assertEqual(by['ZERO']['Data Status'], 'NORMAL')

    def test_15_exceptions_sheet_contains_only_exception_rows_with_source_information(self):
        rows = self.table(load_workbook(self.save())['Exceptions'])
        self.assertEqual([r['Material ID'] for r in rows], ['ALSO MISSING', 'BADQTY', 'BLANK', 'DUP', 'MISSING', 'TEXT', 'TWICE', 'TWICE'])
        self.assertTrue(all(r['Exception Type'] != 'NORMAL' and r['Details'] for r in rows))
        self.assertTrue(all(r['Source File'] == str(self.bom.resolve()) and r['Source Sheet'] == 'DEMO-200-A-40W ' and r['Source Row'] for r in rows))
        by = {r['Material ID']: r for r in rows if r['Material ID'] != 'TWICE'}
        self.assertEqual((by['MISSING']['Assumed / Used Stock'], by['MISSING']['Calculated Shortage'], by['MISSING']['Original Value']), (0, 10, 'Storage: no record'))
        self.assertIn('0+10', by['TEXT']['Original Value']); self.assertIn('BOM QTY: x', by['BADQTY']['Original Value'])
        self.assertIn('row 21', by['DUP']['Details']); self.assertIn('row 22', by['DUP']['Details'])  # both duplicate Storage rows

    def test_16_colour_formatting_red_yellow_none(self):
        sheet = load_workbook(self.save())['Shortage']; result_column = ALL_HEADER.index('Result')
        seen = set()
        for row in sheet.iter_rows(min_row=HEADER_ROW + 1):
            category = row[result_column].value; seen.add(category)
            for cell in row:
                if category == 'ENOUGH':
                    self.assertIsNone(cell.fill.fill_type)
                else:
                    self.assertEqual((cell.fill.fill_type, cell.fill.start_color.rgb[-6:]), ('solid', RED if category == 'SHORTAGE' else YELLOW))
        self.assertEqual(seen, {'SHORTAGE', 'EXCEPTION', 'ENOUGH'})

    def test_17_export_rows_are_sorted_like_the_application(self):
        rows = self.table(load_workbook(self.save())['Shortage'])
        self.assertEqual([(r['Result'], r['Material ID']) for r in rows], [(r.category, r.bom.code) for r in self.checked.sorted_rows])
        self.assertEqual([r['Result'] for r in rows], ['SHORTAGE'] * 3 + ['EXCEPTION'] * 8 + ['ENOUGH'] * 2)

    def test_18_source_workbooks_unchanged_by_check_and_export(self):
        self.save()
        self.assertEqual(self.before, (self.stock.read_bytes(), self.bom.read_bytes()))

    def test_19_export_cannot_overwrite_a_source_workbook_or_any_existing_file(self):
        for target in (self.stock, self.bom, Path(str(self.stock).upper()), self.base / 'sub' / '..' / 'storage.xlsx'):
            with self.subTest(target=target), self.assertRaises(ReportError):
                save_report(target, self.checked, self.stock_book, self.bom_book)
        for name in ('report.csv', 'report.xls', 'report', '~$report.xlsx'):
            with self.subTest(name=name), self.assertRaises(ReportError):
                save_report(self.base / name, self.checked, self.stock_book, self.bom_book)
        self.assertEqual(self.before, (self.stock.read_bytes(), self.bom.read_bytes()))
        other = self.base / 'someone elses workbook.xlsx'; other.write_bytes(b'business data')
        written = save_report(other, self.checked, self.stock_book, self.bom_book)
        self.assertEqual(written.name, 'someone elses workbook (2).xlsx'); self.assertEqual(other.read_bytes(), b'business data')
        self.assertEqual(self.save().name, 'report.xlsx'); self.assertEqual(self.save().name, 'report (2).xlsx'); self.assertEqual(self.save().name, 'report (3).xlsx')

    def test_19b_text_that_looks_like_a_formula_is_written_as_text(self):
        sheet = load_workbook(self.save())['Shortage']
        cell = next(row[1] for row in sheet.iter_rows(min_row=HEADER_ROW + 1) if row[1].value == '=FORMULA')
        self.assertEqual(cell.data_type, 's')  # same text, not a live formula

    def test_13b_default_filename(self):
        when = datetime(2026, 10, 7, 9, 30, 5)
        self.assertEqual(default_filename('DEMO-200-A-40W', 5, when), 'Material_Check_DEMO-200-A-40W_QTY5_20261007_093005.xlsx')
        self.assertEqual(default_filename('A/B:C*?"<>|D', 1, when), 'Material_Check_A_B_C______D_QTY1_20261007_093005.xlsx')
        self.assertEqual(default_filename("DEMO-UNIT 200's", 3, when), "Material_Check_DEMO-UNIT 200's_QTY3_20261007_093005.xlsx")
        self.assertEqual(safe_target(self.base / 'new.xlsx', ()).name, 'new.xlsx')


class UiFlowTests(Files):
    def setUp(self):
        super().setUp()
        self.root = tk.Tk(); self.root.withdraw(); self.app = CheckerApp(self.root)
        self.errors = []
        self.error_patch = patch('checker.ui.messagebox.showerror', side_effect=lambda title, text, **kw: self.errors.append(text))
        self.error_patch.start()

    def tearDown(self):
        self.app.close(); self.error_patch.stop(); super().tearDown()

    def finish(self):
        deadline = time.monotonic() + 10
        while self.app.busy and time.monotonic() < deadline:
            self.root.update(); time.sleep(.005)
        self.assertFalse(self.app.busy, 'Background work did not finish')
        self.root.update()

    def load(self):
        self.app.load_file('stock', str(self.stock)); self.finish()
        self.app.load_file('bom', str(self.bom)); self.finish()
        self.assertFalse(self.errors)

    def check(self, product='DEMO-200-A-40W', quantity='10'):
        self.app.product_text.set(product); self.app.quantity_text.set(quantity)
        self.app.check_button.invoke(); self.finish()

    def shown(self):
        columns = self.app.tree['columns']
        return [dict(zip(columns, self.app.tree.item(i, 'values')), tag=self.app.tree.item(i, 'tags')[0]) for i in self.app.tree.get_children()]

    def test_11_product_id_must_be_selected_before_calculation(self):
        self.load()
        self.assertEqual(tuple(self.app.product_combo['values']), ('DEMO-200-A-40W', 'DEMO-200-A-25W', 'KIT-115'))
        self.check('demo-200')  # two candidates: no calculation, candidates offered
        self.assertIsNone(self.app.result); self.assertIn('2 个匹配项', self.errors[-1])
        self.assertEqual(tuple(self.app.product_combo['values']), ('DEMO-200-A-40W', 'DEMO-200-A-25W'))
        self.check('kit')  # one candidate is still only typed text, not a selection
        self.assertIsNone(self.app.result); self.assertEqual(len(self.errors), 2)
        self.check('NOT-A-PRODUCT')
        self.assertIsNone(self.app.result); self.assertIn('没有匹配', self.errors[-1]); self.assertFalse(self.app.busy)
        self.app.product_text.set('kit'); self.app._pick()  # Enter with one candidate selects it
        self.assertEqual(self.app.product_text.get(), 'KIT-115')
        self.app.product_text.set('demo'); self.app.product_combo.event_generate = lambda *a, **k: None; self.app._pick()
        self.assertEqual(self.app.product_text.get(), 'demo')  # several candidates: nothing chosen for the user
        self.check('demo-200-a-40w')  # full ID typed in another case
        self.assertEqual(len(self.errors), 3); self.assertEqual(self.app.result.product.sheet_name, 'DEMO-200-A-40W ')

    def test_12_production_quantity_input(self):
        self.load()
        for bad in ('', '0', '-3', '1.5', 'ten', '1e2'):
            self.check(quantity=bad); self.assertIsNone(self.app.result)
        self.assertEqual(len(self.errors), 6); self.assertTrue(all('正整数' in e for e in self.errors))
        self.check(quantity='10'); self.assertEqual(self.app.result.production_count, 10)

    def test_result_table_order_colours_columns_and_summary(self):
        self.load(); self.check()
        self.assertEqual(self.app.tree['columns'], ('code', 'description', 'required', 'stock', 'shortage', 'result', 'status', 'rack', 'level'))
        rows = self.shown()
        self.assertEqual([r['result'] for r in rows], ['SHORTAGE'] * 3 + ['EXCEPTION'] * 8 + ['ENOUGH'] * 2)
        self.assertEqual([r['code'] for r in rows], [r.bom.code for r in self.app.result.sorted_rows])
        self.assertTrue(all(r['tag'] == r['result'] for r in rows))
        colour = lambda tag: str(self.app.tree.tag_configure(tag, 'background'))
        self.assertEqual((colour('SHORTAGE'), colour('EXCEPTION'), colour('ENOUGH')), ('#fde9e7', '#fff4d6', 'white'))
        by = {r['code']: r for r in rows if r['code'] != 'TWICE'}
        self.assertEqual((by['SHORT']['required'], by['SHORT']['stock'], by['SHORT']['shortage'], by['SHORT']['status']), ('20', '12', '8', 'NORMAL'))
        self.assertEqual((by['MISSING']['stock'], by['MISSING']['shortage'], by['MISSING']['status'], by['MISSING']['tag']),
                         ('0 (assumed)', '10', 'STOCK_NOT_FOUND / ASSUMED 0', 'EXCEPTION'))  # yellow, not red
        self.assertEqual((by['ZERO']['stock'], by['ZERO']['tag']), ('0', 'SHORTAGE')); self.assertEqual((by['DUP']['stock'], by['DUP']['shortage']), ('—', '—'))
        summary = self.app.summary.get()
        for expected in ('Product ID: DEMO-200-A-40W', 'Production Quantity: 10', 'Total Material Rows: 13', 'SHORTAGE: 3 ',
                         'EXCEPTION: 8 ', 'ENOUGH: 2 ', 'Total calculated shortage rows: 6'):
            self.assertIn(expected, summary)
        self.assertNotIn('READY', summary.upper()); self.assertNotIn('APPROV', summary.upper())
        target = next(i for i in self.app.tree.get_children() if self.app.tree.item(i, 'values')[0] == 'TEXT')
        self.app.tree.selection_set(target); self.app.show_details(); details = self.app.details.get('1.0', 'end')
        self.assertIn('0+10', details); self.assertIn(str(self.bom.resolve()), details); self.assertIn("'DEMO-200-A-40W '", details)

    def test_export_excel_from_the_ui_uses_suggested_name_and_never_overwrites(self):
        self.load(); self.check(); before = (self.stock.read_bytes(), self.bom.read_bytes()); asked = {}
        def dialog(**options):
            asked.update(options); return str(self.base / options['initialfile'])
        with patch('checker.ui.filedialog.asksaveasfilename', side_effect=dialog):
            self.app.export_button.invoke(); self.app.export_button.invoke()
        self.assertRegex(asked['initialfile'], r'^Material_Check_DEMO-200-A-40W_QTY10_\d{8}_\d{6}\.xlsx$'); self.assertFalse(asked['confirmoverwrite'])
        reports = sorted(p.name for p in self.base.glob('Material_Check_*.xlsx'))
        self.assertEqual(len(reports), 2); self.assertFalse(self.errors)  # second export got its own name
        self.assertEqual(load_workbook(self.base / reports[0]).sheetnames, ['Shortage', 'Exceptions'])
        self.assertIsNone(self.app.write_export(str(self.stock))); self.assertEqual(len(self.errors), 1)
        self.assertEqual(before, (self.stock.read_bytes(), self.bom.read_bytes()))

    def test_20_file_update_refreshes_result_and_failure_shows_no_stale_result(self):
        self.load(); self.check()
        self.assertEqual(self.app.result.category_count('SHORTAGE'), 3); parses = self.app.service.cache.parse_count
        self.check(); self.assertEqual(self.app.service.cache.parse_count, parses)  # unchanged files: cache
        workbook(self.stock, [('Stock', stock_rows([(c, 1000 if c == 'SHORT' else q) for c, q in STOCK]))])
        self.check()
        self.assertEqual(self.app.service.cache.parse_count, parses + 1); self.assertEqual(self.app.result.category_count('SHORTAGE'), 2)
        self.assertIn('已从文件重新读取', self.app.file_status['stock'].get())
        os.replace(self.stock, self.base / 'away.bin'); self.check()
        self.assertEqual(len(self.errors), 1); self.assertIsNone(self.app.result); self.assertEqual(self.app.tree.get_children(), ())
        self.assertIsNotNone(self.app.bom_book); self.assertEqual(self.app.product_text.get(), 'DEMO-200-A-40W')
        self.assertIsNone(self.app.write_export(str(self.base / 'stale.xlsx'))); self.assertFalse((self.base / 'stale.xlsx').exists())


@unittest.skipUnless(os.environ.get('MRC_TEST_MOUNT_DIR'), 'set MRC_TEST_MOUNT_DIR to a mounted folder (e.g. Google Drive root) to run')
class DriveMountTests(Files):
    """21: synthetic test files on a real mount. Nothing outside the temporary sub-folder is touched."""
    folder = os.environ.get('MRC_TEST_MOUNT_DIR')
    quantity = staticmethod(lambda book: book.stock_sheets[0].materials[0].quantity)

    def write_stock(self, short, path=None):
        return workbook(path or self.stock, [('Stock', stock_rows([('SHORT', short)] + STOCK[1:]))])

    def test_read_cache_refresh_and_identical_size_and_time(self):
        first = self.cache.read(str(self.stock), 'stock')
        self.assertIs(first, self.cache.read(str(self.stock), 'stock')); self.assertEqual(self.cache.parse_count, 1)
        stat = self.stock.stat(); self.write_stock(13); os.utime(self.stock, ns=(stat.st_atime_ns, stat.st_mtime_ns))
        self.assertEqual(self.stock.stat().st_size, stat.st_size)
        second = self.cache.read(str(self.stock), 'stock')
        self.assertIsNot(first, second); self.assertEqual(self.quantity(second), 13)

    def test_replacement_and_temporary_unavailability(self):
        self.cache.read(str(self.stock), 'stock')
        os.replace(self.write_stock(14, self.base / 'incoming.tmp'), self.stock)
        self.assertEqual(self.quantity(self.cache.read(str(self.stock), 'stock')), 14)
        os.replace(self.stock, self.base / 'away.bin')
        with self.assertRaises(DataReadError):
            self.cache.read(str(self.stock), 'stock')
        os.replace(self.base / 'away.bin', self.stock)
        self.assertEqual(self.quantity(self.cache.read(str(self.stock), 'stock')), 14)

    @unittest.skipUnless(os.name == 'nt', 'Windows share modes')
    def test_locked_workbook(self):
        kernel = ctypes.WinDLL('kernel32', use_last_error=True)
        kernel.CreateFileW.argtypes = [ctypes.c_wchar_p, ctypes.c_uint32, ctypes.c_uint32, ctypes.c_void_p, ctypes.c_uint32, ctypes.c_uint32, ctypes.c_void_p]
        kernel.CreateFileW.restype = ctypes.c_void_p; kernel.CloseHandle.argtypes = [ctypes.c_void_p]
        for access, share, readable in ((0x80000000, 0, False), (0xC0000000, 1, True)):  # exclusive; editor-style shared read
            handle = kernel.CreateFileW(str(self.stock), access, share, None, 3, 0x80, None)
            self.assertNotEqual(handle, ctypes.c_void_p(-1).value)
            try:
                if readable:
                    self.assertEqual(self.quantity(ReadOnlyCache().read(str(self.stock), 'stock')), 12)
                else:
                    with self.assertRaises(DataReadError):
                        ReadOnlyCache().read(str(self.stock), 'stock')
            finally:
                kernel.CloseHandle(handle)

    def test_source_change_during_read_is_rejected_then_recovers(self):
        real_load, real_fingerprint, state = reader.load_workbook, reader.fingerprint, {}
        def load(stream, **kwargs):
            book = real_load(stream, **kwargs); state['new'] = self.write_stock(15, self.base / 'during.tmp'); return book
        def fingerprint(path):
            if state.get('new') and state['new'].exists():
                os.replace(state['new'], self.stock)
            return real_fingerprint(path)
        with patch.object(reader, 'load_workbook', load), patch.object(reader, 'fingerprint', fingerprint):
            with self.assertRaisesRegex(DataReadError, '发生变化'):
                self.cache.read(str(self.stock), 'stock')
        self.assertEqual(self.quantity(self.cache.read(str(self.stock), 'stock')), 15)

    def test_owner_file_m2_check_and_report_on_the_mount(self):
        (self.base / '~$storage.xlsx').write_bytes(b'owner')
        service = CheckerService(); stock, locks = service.load(str(self.stock), 'stock'); bom, _ = service.load(str(self.bom), 'bom')
        self.assertEqual(locks, ('~$storage.xlsx',))
        before = (self.stock.read_bytes(), self.bom.read_bytes())
        result = check_materials(bom.products[0], stock.stock_sheets[0], 10)
        self.assertEqual([result.category_count(c) for c in ('SHORTAGE', 'EXCEPTION', 'ENOUGH')], [3, 8, 2])
        written = save_report(self.base / 'report.xlsx', result, stock, bom)
        book = load_workbook(written)
        self.assertEqual(book.sheetnames, ['Shortage', 'Exceptions']); self.assertEqual(book['Shortage'].max_row, HEADER_ROW + len(BOM))
        with self.assertRaises(ReportError):
            save_report(self.stock, result, stock, bom)
        self.assertEqual(before, (self.stock.read_bytes(), self.bom.read_bytes()))


if __name__ == '__main__':
    unittest.main()
