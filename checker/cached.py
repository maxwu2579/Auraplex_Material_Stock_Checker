"""Explicit, session-only historical SQLite calculations. Never verifies Excel or cloud freshness."""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
import hashlib
import json
from pathlib import Path
import sqlite3

from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font

from .core import WorkbookData, ROW_CAVEAT
from .inventory_db import SUPPORTED_SCHEMA_VERSIONS
from .queries import open_readonly, current_version, load_workbook_data
from . import report

LABEL = "CACHED DATA / UNVERIFIED"


class CacheError(ValueError):
    """Historical data is missing, unreadable, changed, or not explicitly selected."""


class CacheReadError(CacheError):
    """The database itself could not be read safely; previous status is no longer usable."""


@dataclass(frozen=True)
class Version:
    role: str
    id: int
    path: str
    imported_at: str
    source_modified_at: str
    last_verified_at: str | None
    content_signature: str
    file_sha256: str


@dataclass(frozen=True)
class Snapshot:
    database: str
    identity: tuple
    versions: tuple[Version, ...]
    stock: WorkbookData
    bom: WorkbookData
    _content: str = field(init=False, repr=False)

    def __post_init__(self):
        object.__setattr__(self, '_content', hashlib.sha256(repr((self.stock, self.bom)).encode('utf-8')).hexdigest())

    @property
    def token(self):
        # Includes actual reconstructed rows, not just IDs which a replaced database could reuse.
        return (self.database, self.identity, tuple((v.role, v.id, v.path, v.imported_at,
                v.content_signature, v.file_sha256) for v in self.versions), self._content)

    def describe(self):
        lines = [LABEL, "Cloud Not Verified. Excel files will NOT be checked.",
                 "Storage and BOM are independent historical versions; their import times may differ."]
        for v in self.versions:
            lines.append(f"\n{'Storage' if v.role == 'stock' else 'BOM'} · version {v.id}\n{v.path}\n"
                         f"Source modified (UTC): {v.source_modified_at}\nImported (UTC): {v.imported_at}\n"
                         f"Previously verified (UTC): {v.last_verified_at or 'Unknown'}")
        return '\n'.join(lines)


@dataclass(frozen=True)
class CacheBinding:
    snapshot: Snapshot
    checked_at: float
    valid: bool = True

    @property
    def current(self):
        return False  # A database observation can NEVER become local-source verification.

    def describe(self):
        return (self.snapshot.describe() +
                f"\nCalculated at {datetime.fromtimestamp(self.checked_at).astimezone().isoformat()}\n"
                + ("Historical versions selected explicitly." if self.valid else "OUTDATED / RECHECK REQUIRED"))


def load(database) -> Snapshot:
    """Read both roles in one SQLite snapshot; no creation, migration, source access or writes."""
    conn = None
    try:
        path = Path(database).resolve()
        before = path.stat()
        conn = open_readonly(path)
        conn.execute('PRAGMA busy_timeout = 1500')
        conn.execute('BEGIN')
        if conn.execute('PRAGMA user_version').fetchone()[0] not in SUPPORTED_SCHEMA_VERSIONS:
            raise CacheReadError('Unsupported cache database schema; historical calculation refused.')
        versions, books = [], {}
        for role in ('stock', 'bom'):
            row = current_version(conn, role)
            if row is None:
                raise CacheError(f'No successful {role} cache. Both Storage and BOM are required.')
            success = conn.execute("SELECT max(finished_at) FROM sync_runs WHERE role = ? "
                                   "AND source_version_id = ? AND outcome IN ('IMPORTED', 'UNCHANGED')",
                                   (role, row['id'])).fetchone()[0]
            if not success:
                raise CacheError(f'No successful synchronisation record for cached {role}.')
            versions.append(Version(role, row['id'], row['source_path'], row['imported_at'],
                                    row['source_modified_at'], success, row['content_signature'], row['file_sha256']))
            books[role] = load_workbook_data(conn, role, row['id'])
            sheets = books[role].stock_sheets if role == 'stock' else books[role].products
            if len(sheets) != row['sheet_count'] or sum(len(s.materials) for s in sheets) != row['record_count']:
                raise CacheReadError(f'Cached {role} rows are incomplete; historical calculation refused.')
        if not books['stock'].stock_sheets or not books['bom'].products:
            raise CacheError('Cached workbook structure is incomplete.')
        after = path.stat()
        identity = (before.st_dev, before.st_ino)
        if identity != (after.st_dev, after.st_ino):
            raise CacheError('Cache database was replaced while reading. Please select it again.')
        return Snapshot(str(path), identity, tuple(versions), books['stock'], books['bom'])
    except (OSError, sqlite3.Error, LookupError, ValueError, TypeError, ArithmeticError) as exc:
        if isinstance(exc, CacheError):
            raise
        raise CacheReadError(f'Cache database cannot be read: {exc}') from exc
    finally:
        if conn is not None:
            conn.close()


def verify(snapshot):
    if not isinstance(snapshot, Snapshot) or load(snapshot.database).token != snapshot.token:
        raise CacheError('Cached Storage/BOM versions changed. Select cached data again and CHECK.')


def save(path, result, stock, bom, bound):
    """Separate historical export contract. Cannot substitute for the verified Excel writer."""
    if not isinstance(bound, CacheBinding) or not bound.valid:
        raise CacheError('No valid explicit cached CHECK. Please select cached data and CHECK.')
    snapshot = bound.snapshot
    if stock != snapshot.stock or bom != snapshot.bom or result.product not in bom.products:
        raise CacheError('Report data does not belong to the selected cached versions.')
    verify(snapshot)
    target = Path(path)
    if not target.name.startswith('UNVERIFIED_'):
        # Validate the original choice first: a source path must be refused, never silently renamed.
        report.safe_target(target, (stock.path, bom.path))
        target = target.with_name('UNVERIFIED_' + target.name)

    def build():
        book = report.build_report(result, stock, bom, datetime.now())
        for sheet in book:
            sheet['A1'] = 'UNVERIFIED · Cached Material Requirement & Stock Report'
            sheet['A1'].font = Font(name='Segoe UI', size=16, bold=True, color='A53432')
            sheet.merge_cells(start_row=8, start_column=1, end_row=8, end_column=sheet.max_column)
            sheet['A8'] = LABEL + ' · Cloud Not Verified · Historical SQLite data; Excel sources not checked.'
            sheet['A8'].alignment = Alignment(wrap_text=True, vertical='center')
            sheet['A8'].font = Font(bold=True, color='A53432')
            sheet.row_dimensions[8].height = 32
            sheet.print_title_rows = '8:15'  # The UNVERIFIED warning repeats on every printed page.
            sheet['A9'] = 'Calculated from historical SQLite snapshots. Verification Required. ' + ROW_CAVEAT
            for v, row in zip(snapshot.versions, (5, 6)):
                cell = sheet.cell(row, 3)
                cell.value = report._value(f'{v.path}\nCached version {v.id} · Imported UTC: {v.imported_at}')
                cell.data_type = 's'
                sheet.row_dimensions[row].height = 44
                sheet.cell(row, 3).comment = Comment(json.dumps({**v.__dict__, 'mode': LABEL,
                    'calculated_at': datetime.fromtimestamp(bound.checked_at).astimezone().isoformat(),
                    'scope': 'Historical data only; Excel and cloud freshness UNVERIFIED'},
                    ensure_ascii=False, indent=2), 'Material Stock Checker')
        return book

    return report._stage_report(target, (stock.path, bom.path), build, lambda: verify(snapshot))
