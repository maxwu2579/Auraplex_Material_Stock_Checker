"""Writes a check result to a NEW .xlsx report. Source workbooks are never opened for writing."""
from __future__ import annotations

from datetime import datetime
from decimal import Decimal
from contextlib import contextmanager, nullcontext
from copy import copy
import hashlib
import json
import os
from pathlib import Path
import re
import tempfile
import threading

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.comments import Comment
from openpyxl.styles import Font

from .core import CheckResult, EXCEPTION, ROW_CAVEAT, SHORTAGE, WorkbookData
from .report_style import FILLS, HEADER_ROW, METADATA_START, NOTE_ROW, format_report
from . import binding
from .sync import content_signature


class ReportError(ValueError):
    """The chosen report destination is not acceptable; nothing was written."""


class ExportCancelled(Exception):
    """The export was cancelled before publication: no report was published. Its own staging file is removed; if that
    removal fails, `cleanup_warning` says which file remains and why."""


class Cancellation:
    """Cancellation of one export, serialised with its publication.

    Cancelling and publishing take the same lock, so exactly one of them happens first. If cancel() wins, the export
    is never published. If the publication wins, the report stays published and cancel() reports that. The lock is
    held only for the cancellation check and the publishing rename/link, never while a report is generated, written
    or verified, so the caller of cancel() (the Tk thread) waits at most for one filesystem rename."""

    def __init__(self):
        self._lock = threading.Lock()
        self._cancelled = False
        self.published = None    # the published report path, once publication has won

    def is_set(self):
        return self._cancelled

    def cancel(self):
        """Returns True when the cancellation came first (nothing was or will be published)."""
        with self._lock:
            self._cancelled = True
            return self.published is None

    @contextmanager
    def publishing(self, target):
        with self._lock:
            if self._cancelled:
                raise ExportCancelled("Export cancelled before the report was published.")
            yield
            self.published = target


# Stages an export reports through its optional `progress` callback (no percentage is invented for any of them).
CHECKING, BUILDING, WRITING, VERIFYING = "checking", "building", "writing", "verifying"


def _step(stage, progress, cancelled):
    """Cancellation boundary before each stage; a stage already running (e.g. the workbook save) is never killed."""
    if cancelled is not None and cancelled.is_set():
        raise ExportCancelled("Export cancelled before the report was published.")
    if progress is not None and stage is not None:
        progress(stage)


ALL_HEADER = ("No.", "Material ID", "Description", "BOM QTY", "Production Quantity", "Required", "Stock", "Shortage",
              "Result", "Data Status", "Rack", "Level", "Original Stock Value", "Source Sheet", "Source Row")
EXCEPTION_HEADER = ("No.", "Material ID", "Description", "Required", "Assumed / Used Stock", "Calculated Shortage",
                    "Exception Type", "Original Value", "Details", "Rack", "Level", "Source File", "Source Sheet", "Source Row")
NOTE = ("Calculated from the files as last saved. Not a production approval. Rows marked EXCEPTION need a person to check the "
        "source data; where Data Status says ASSUMED 0, the stock value 0 was assumed, not read from Storage.")


def default_filename(product_id: str, quantity: int, when: datetime) -> str:
    # Only characters Windows forbids in file names are replaced; the Product ID shown elsewhere is unchanged.
    safe = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", product_id).strip(" .") or "product"
    return f"Material_Check_{safe}_QTY{quantity}_{when:%Y%m%d_%H%M%S}.xlsx"


def _same_file(a: Path, b: Path) -> bool:
    try:
        if a.exists() and b.exists() and os.path.samefile(a, b):
            return True
    except OSError:
        pass
    return os.path.normcase(os.path.abspath(a)) == os.path.normcase(os.path.abspath(b))


def safe_target(path: str | Path, sources) -> Path:
    """A path that is a .xlsx, is not a source workbook, and does not exist yet (a number is added if needed)."""
    target = Path(path)
    if target.suffix.lower() != ".xlsx":
        raise ReportError("报告只能保存为新的 .xlsx 文件。")
    if target.name.startswith("~$"):
        raise ReportError("不能使用 ~$ 开头的 Excel 占用标记文件名。")
    if any(_same_file(target, Path(source)) for source in sources):
        raise ReportError("导出目标是当前的 Storage 或 BOM 源工作簿；已拒绝，源文件不会被覆盖。")
    candidate, n = target, 1
    while candidate.exists():  # never overwrite anything, including unrelated workbooks
        n += 1
        candidate = target.with_name(f"{target.stem} ({n}){target.suffix}")
    return candidate


def _value(value):
    """Cell value: numbers stay numbers, everything else becomes text that Excel cannot run."""
    if value is None or isinstance(value, bool):
        return None if value is None else str(value)
    if isinstance(value, Decimal):
        return int(value) if value == value.to_integral_value() else float(value)
    if isinstance(value, (int, float)):
        return value
    return ILLEGAL_CHARACTERS_RE.sub("", str(value))


def _write_row(sheet, row, values, fill=None, bold=False):
    for column, value in enumerate(values, 1):
        cell = sheet.cell(row=row, column=column, value=_value(value))
        if isinstance(cell.value, str) and cell.data_type == "f":
            cell.data_type = "s"  # text copied from a source workbook must never become a live formula
        if fill is not None:
            cell.fill = fill
        if bold:
            cell.font = Font(bold=True)


def _joined(values):
    return " | ".join("" if v is None else str(v) for v in values)


def _original_value(r):
    parts = []
    if "INVALID_MATERIAL_ID" in r.exceptions or "CODE_TYPE_MISMATCH" in r.exceptions:
        parts.append(f"BOM Material ID: {r.bom.code!r}")
    if "INVALID_BOM_QTY" in r.exceptions:
        parts.append(f"BOM QTY: {'(blank)' if r.bom.quantity is None else r.bom.quantity}")
    if r.stock:
        parts.append(f"Storage LIVE STOCK: {r.original_stock}")
    elif "STOCK_NOT_FOUND" in r.exceptions:
        parts.append("Storage: no record")
    return "; ".join(parts)


def _details(r):
    where = _joined(f"{s.source.sheet}!row {s.source.row} (code {s.code!r})" for s in r.stock)
    return " ".join(r.issues) + (f" Storage records: {where}." if where else " Storage records: none.") + \
        (f" Assembly section: {r.bom.section}." if r.bom.section else "")


def build_report(result: CheckResult, stock_book: WorkbookData, bom_book: WorkbookData, generated: datetime) -> Workbook:
    book = Workbook()
    rows = result.sorted_rows
    for sheet, title in ((book.active, "Shortage"), (book.create_sheet(), "Exceptions")):
        sheet.title = title
        for row, (label, value) in enumerate((("Product ID", result.product.display_name),
                                              ("Production Quantity", result.production_count),
                                              ("Storage File", stock_book.path), ("BOM File", bom_book.path),
                                              ("Generated Time", generated.strftime("%Y-%m-%d %H:%M:%S"))), METADATA_START):
            _write_row(sheet, row, (label, None, value))
            sheet.cell(row=row, column=1).font = Font(bold=True)
        _write_row(sheet, NOTE_ROW, (NOTE + " " + ROW_CAVEAT,))
    everything, exceptions = book["Shortage"], book["Exceptions"]
    _write_row(everything, HEADER_ROW, ALL_HEADER, bold=True)
    for n, r in enumerate(rows, 1):
        _write_row(everything, HEADER_ROW + n,
                   (n, r.bom.code, r.bom.description, r.bom.quantity, result.production_count, r.required, r.stock_used,
                    r.shortage, r.category, r.data_status, _joined(s.rack for s in r.stock), _joined(s.level for s in r.stock),
                    r.original_stock, r.bom.source.sheet, r.bom.source.row), fill=FILLS.get(r.category))
    _write_row(exceptions, HEADER_ROW, EXCEPTION_HEADER, bold=True)
    for n, r in enumerate((r for r in rows if r.category == EXCEPTION), 1):
        _write_row(exceptions, HEADER_ROW + n,
                   (n, r.bom.code, r.bom.description, r.required, r.stock_used, r.shortage, r.data_status, _original_value(r),
                    _details(r), _joined(s.rack for s in r.stock), _joined(s.level for s in r.stock), r.bom.source.path,
                    r.bom.source.sheet, r.bom.source.row))
    format_report(book, result, stock_book, bom_book, _write_row)
    return book


def save_report(path: str | Path, result: CheckResult, stock_book: WorkbookData, bom_book: WorkbookData,
                generated: datetime | None = None) -> Path:
    """Write the report to a new file and return the path actually used. Raises ReportError before writing anything."""
    target = safe_target(path, (stock_book.path, bom_book.path))
    build_report(result, stock_book, bom_book, generated or datetime.now()).save(target)
    return target


def _export_provenance(book, checked, verified):
    """Keep the two M2 tables intact; provenance occupies the unused row 8 and source-cell comments."""
    checked_at = datetime.fromtimestamp(checked.checked_at).astimezone().isoformat(timespec="microseconds")
    verified_at = datetime.fromtimestamp(verified.checked_at).astimezone().isoformat(timespec="microseconds")
    note = (f"Local versions checked for export: {verified_at} | Cloud Not Verified. "
            "Point-in-time local check only; sources can change afterwards.")
    for sheet in book:
        sheet.merge_cells(start_row=8, start_column=1, end_row=8, end_column=sheet.max_column)
        _write_row(sheet, 8, (note,))
        sheet['A8'].alignment = copy(sheet['A9'].alignment)
        sheet.row_dimensions[8].height = 30
        for role, row in (("stock", 5), ("bom", 6)):
            version = checked.versions[role]
            metadata = {"source": binding.ROLE_NAMES[role], "path": version[0], "size_bytes": version[1],
                        "mtime_ns": version[2], "ctime_ns": version[3], "device": version[4], "inode": version[5],
                        "content_signature": content_signature(version),
                        "version_id": hashlib.sha256(json.dumps(version, ensure_ascii=False).encode('utf-8')).hexdigest(),
                        "check_verified_at": checked_at, "export_verified_at": verified_at,
                        "scope": "Local file version at verification time only; Cloud Not Verified"}
            sheet.cell(row, 3).comment = Comment(json.dumps(metadata, ensure_ascii=False, indent=2), "Material Stock Checker")


def save_verified_report(path, result, stock_book, bom_book, source_binding, source_paths, *, progress=None,
                         cancelled=None):
    """Application export boundary: stage privately, recheck, then publish a new file without overwriting.

    No retries or SQLite reads; failures leave no final report. The low-level save_report remains available for
    formatting raw calculation data, whereas the application's exports must use this version-bound boundary.
    progress(stage) and cancelled (an Event) are optional; see _stage_report for the cancellation boundaries.
    """
    paths = dict(source_paths)
    checked = source_binding
    _step(CHECKING, progress, cancelled)
    verified = binding.verify_export(checked, paths)
    for role, source in (("stock", stock_book), ("bom", bom_book)):
        if source.fingerprint != checked.versions[role]:
            confirmed = dict(verified.confirmed)
            confirmed[role] = None
            failed = binding.SourceBinding(dict(checked.versions), confirmed, verified.via_sync,
                                           verified.checked_at, checked.attempts)
            raise binding.ExportVerificationError(failed, f"{binding.ROLE_NAMES[role]} report data differs from CHECK.")
    sources = (stock_book.path, bom_book.path)
    def build():
        book = build_report(result, stock_book, bom_book, datetime.now())
        try:
            _export_provenance(book, checked, verified)
            return book
        except BaseException:
            book.close()
            raise
    return _stage_report(path, sources, build, lambda: binding.verify_export(checked, paths), progress=progress,
                         cancelled=cancelled)


def _stage_report(path, sources, build, validate, *, progress=None, cancelled=None):
    """Shared no-overwrite publication; caller supplies its own explicit verification contract.

    Cancellation is checked before building, before writing and before the final verification; these steps and the
    verification itself run to completion. With a Cancellation, the last check and the publication are serialised
    with cancel(): a cancellation that comes first prevents publication, one that comes later never removes the
    published report. A cancelled export removes only its own staging file (a removal failure is reported through
    `cleanup_warning` on the exception, never hidden)."""
    _step(None, None, cancelled)
    target = safe_target(path, sources)
    descriptor, temporary = tempfile.mkstemp(prefix=".mrc-export-", suffix=".xlsx", dir=target.parent)
    os.close(descriptor)
    failure = None
    try:
        _step(BUILDING, progress, cancelled)
        book = build()
        try:
            _step(WRITING, progress, cancelled)
            book.save(temporary)
        finally:
            book.close()
        _step(VERIFYING, progress, cancelled)
        # This last observation is after report generation. The report conservatively records the earlier
        # completed export verification, not a future timestamp that was not known when the workbook was saved.
        for _ in range(10):
            target = safe_target(path, sources)
            validate()
            # The publication boundary: the cancellation check and the rename/link are one step under the
            # Cancellation lock (a plain event only gets the check). Cancelled first: nothing is published.
            guard = cancelled.publishing(target) if hasattr(cancelled, "publishing") else nullcontext()
            _step(None, None, cancelled)
            try:
                with guard:
                    if os.name == "nt":
                        os.rename(temporary, target)  # Windows refuses an existing destination atomically
                    else:
                        os.link(temporary, target)    # same no-overwrite contract on POSIX
                return target
            except FileExistsError:
                continue
        raise ReportError("报告目标持续被其他程序创建；未发布报告，请选择其他名称。")
    except BaseException as exc:
        failure = exc
        raise
    finally:
        try:
            os.unlink(temporary)  # only our own staging file; never an Excel source or another report
        except FileNotFoundError:
            pass  # Windows publication already consumed it, or another cleanup removed it
        except OSError as cleanup:
            if failure is None:
                raise
            # Cleanup must never replace a source-verification failure (or the original write error).
            # The UI can still invalidate the binding and show where a staging file may remain.
            failure.cleanup_warning = f"Temporary report cleanup failed: {temporary}\n{cleanup}"
