from __future__ import annotations

from collections import Counter
from dataclasses import replace
from pathlib import Path
from threading import Event
import time
from zipfile import BadZipFile, ZipFile
import xml.etree.ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .core import DataReadError, Material, Product, Source, StockSheet, WorkbookData, normalize_code, number, material_key


STOCK_HEADER = {5: "PART NO", 6: "DESCRIPTION", 11: "RACKING", 12: "LEVEL", 15: "LIVE STOCK"}
BOM_HEADER = {4: "PART NO", 5: "DESCRIPTION", 12: "QTY"}
HEADER_SCAN_ROWS = 100


def fingerprint(path: str | Path) -> tuple:
    p = Path(path).expanduser()
    try:
        p = p.resolve(strict=True)
        if not p.is_file():
            raise DataReadError("请选择 Excel 文件，而不是文件夹。")
        if p.name.startswith("~$"):
            raise DataReadError("不能选择 ~$ Excel 临时锁文件。")
        if p.suffix.lower() not in (".xlsx", ".xlsm"):
            raise DataReadError("M0 只支持 .xlsx / .xlsm；.xls / .csv 尚未支持。")
        s = p.stat()
        # Central-directory CRCs detect changed data even when a sync preserves size/mtime.
        # No media bytes are decompressed/read. This is an invalidation signature, not a security hash.
        with ZipFile(p, "r") as archive:
            signature = tuple(sorted((i.filename, i.CRC, i.file_size, i.compress_size) for i in archive.infolist()
                                     if i.filename.startswith("xl/") and not i.filename.startswith(("xl/media/", "xl/drawings/"))))
        return (str(p), s.st_size, s.st_mtime_ns, s.st_ctime_ns, s.st_dev, s.st_ino, signature)
    except (OSError, RuntimeError, BadZipFile) as exc:
        raise DataReadError(f"文件不存在、不可访问或路径无效：{p}（{exc}）") from exc


def owner_lock_files(path: str | Path) -> tuple[str, ...]:
    """Names of Excel owner files (~$...) beside the workbook. Presence means the workbook is, or
    was, open in Excel somewhere; it does not show whether unsaved changes exist."""
    p = Path(path)
    # Excel prefixes the name with ~$; for long names it overwrites the first two characters instead.
    names = {"~$" + p.name, "~$" + p.name[2:]}
    try:
        return tuple(sorted(n for n in names if (p.parent / n).is_file()))
    except OSError:
        return ()


class ReadOnlyCache:
    """At most one snapshot per selected file/role, all in memory; never writes Excel."""

    def __init__(self):
        self._cache: dict[tuple[str, str], WorkbookData] = {}
        self.parse_count = 0
        self.cache_hits = 0

    def read(self, path: str, role: str, force: bool = False, cancelled: Event | None = None) -> WorkbookData:
        if role not in ("stock", "bom"):
            raise ValueError("Unknown workbook role")
        before = fingerprint(path)
        key = (before[0], role)
        old = self._cache.get(key)
        if not force and old is not None and old.fingerprint == before:
            self.cache_hits += 1
            return old
        self._cache.pop(key, None)
        try:
            # Explicit rb handle + read_only prevents writes. No images, VBA or external links loaded.
            with open(before[0], "rb") as stream:
                workbook = load_workbook(stream, read_only=True, data_only=False, keep_links=False, keep_vba=False)
                try:
                    self.parse_count += 1
                    products, stock_sheets, notes = [], [], []
                    for i, sheet in enumerate(workbook.worksheets):
                        if cancelled is not None and cancelled.is_set():
                            raise DataReadError("读取已取消。")
                        if "TEMPLATE" in normalize_code(sheet.title) or "模板" in sheet.title:
                            notes.append(f"{sheet.title!r}：模板，未作为产品/库存。")
                            continue
                        header = STOCK_HEADER if role == "stock" else BOM_HEADER
                        header_row = None
                        materials = []
                        section = ""
                        # Ignore inflated dimensions; parse actual worksheet rows, max 18 columns.
                        sheet.reset_dimensions()
                        for cells in sheet.iter_rows(max_col=18):
                            if cancelled is not None and cancelled.is_set():
                                raise DataReadError("读取已取消。")
                            populated = [c for c in cells if c.value is not None]
                            if not populated:
                                continue
                            row_no = populated[0].row
                            if header_row is None and row_no > HEADER_SCAN_ROWS:
                                break
                            value = lambda c: cells[c - 1].value
                            if normalize_code(value(2)) == "ASSY NAME:":
                                section = str(value(5) or "")
                                continue
                            is_header = all(normalize_code(value(c)) == label for c, label in header.items())
                            if is_header:
                                if header_row is None:
                                    header_row = row_no
                                continue
                            if header_row is None:
                                continue
                            code_col, description_col, qty_col = (5, 6, 15) if role == "stock" else (4, 5, 12)
                            code_cell, qty_cell = cells[code_col - 1], cells[qty_col - 1]
                            code, description, qty = code_cell.value, value(description_col), qty_cell.value
                            if role == "stock":
                                # Exclude area headings such as SCREW STORE, not incomplete material rows.
                                area_heading = (description is None and qty is None and value(11) is None and value(12) is None
                                                and normalize_code(value(1)) in ("GROUND FLOOR", "MEZZANITE FLOOR", "MEZZANINE FLOOR"))
                                is_material = (code not in (None, "") and not area_heading) or (
                                    code in (None, "") and description not in (None, "") and
                                    (number(value(1)) is not None or any(value(c) is not None for c in (11, 12, 15))))
                            else:
                                is_material = code not in (None, "") or (
                                    number(value(2)) is not None and (description not in (None, "") or qty is not None))
                            if not is_material:
                                continue
                            materials.append(Material(Source(before[0], sheet.title, row_no), code, description, qty,
                                                      qty_cell.data_type, code_cell.data_type,
                                                      value(11) if role == "stock" else value(10),
                                                      value(12) if role == "stock" else value(11), section))
                        if header_row is None:
                            notes.append(f"{sheet.title!r}：未找到 M0 所需表头组合（前100行），未读取为{role}。")
                        elif not materials:
                            notes.append(f"{sheet.title!r}：无材料数据，未作为产品/库存。")
                        elif role == "stock":
                            stock_sheets.append(StockSheet(sheet.title, header_row, tuple(materials)))
                            invalid_codes = sum(not material_key(m) for m in materials)
                            if invalid_codes:
                                notes.append(f"{sheet.title!r}：{invalid_codes}条库存编号为空、非文本、公式或错误，不用于连接。")
                        else:
                            products.append(Product(before[0], sheet.title, f"{sheet.title.strip()}  [Sheet {i + 1}]",
                                                    header_row, tuple(materials)))
                    # Product ID is the trimmed Sheet name; the Sheet number is kept only to separate look-alike names.
                    names = Counter(normalize_code(p.sheet_name) for p in products)
                    products = [p if names[normalize_code(p.sheet_name)] > 1 else replace(p, display_name=p.sheet_name.strip())
                                for p in products]
                    if role == "stock" and not stock_sheets:
                        raise DataReadError("未找到 F01 库存结构：E PART NO、F DESCRIPTION、K RACKING、L LEVEL、O LIVE STOCK。\n"
                                            "旧版 N/L 库存列不支持，不能直接当成 F01 使用。")
                    if role == "bom" and not products:
                        raise DataReadError("未找到含材料的产品 Sheet：D PART NO、E DESCRIPTION、L QTY。模板与空表不参与。")
                    result = WorkbookData(before[0], before, tuple(products), tuple(stock_sheets), tuple(notes), time.time())
                finally:
                    workbook.close()
            if fingerprint(before[0]) != before:
                raise DataReadError("读取期间源文件发生变化，请等待文件保存/同步完成后重新读取。")
        except DataReadError:
            raise
        except (OSError, BadZipFile, InvalidFileException, ET.ParseError, ValueError, KeyError) as exc:
            raise DataReadError(f"无法只读读取文件。可能被独占、权限不足、加密或格式损坏。\n{before[0]}\n{exc}") from exc
        # Evict previous selections for this role to bound memory use.
        self._cache = {k: val for k, val in self._cache.items() if k[1] != role}
        self._cache[key] = result
        return result
