"""Formatting for the existing two report tables. Does not calculate business values."""
from collections import Counter
from pathlib import Path
import math
import unicodedata

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .core import ENOUGH, EXCEPTION, SHORTAGE

HEADER_ROW, METADATA_START, METADATA_VALUE_COLUMN, NOTE_ROW = 15, 3, 3, 9
RED, YELLOW = "FDE9E7", "FFF4D6"
INK, MUTED, DARK_RED, DARK_AMBER = "243244", "687789", "A53432", "95641A"
FILLS = {SHORTAGE: PatternFill("solid", fgColor=RED), EXCEPTION: PatternFill("solid", fgColor=YELLOW)}


def format_report(book, result, stock_book, bom_book, write_row):
    edge = Side(style="thin", color="E2E7EC")
    rows = result.sorted_rows
    numeric = {"No.", "BOM QTY", "Production Quantity", "Required", "Stock", "Shortage",
               "Assumed / Used Stock", "Calculated Shortage", "Source Row"}
    limits = {"No.": (7, 7), "Material ID": (24, 48), "Description": (42, 58),
              "Data Status": (28, 46), "Exception Type": (28, 46), "Original Value": (24, 42),
              "Original Stock Value": (23, 38), "Details": (48, 72), "Source File": (40, 60),
              "Source Sheet": (26, 42), "Rack": (10, 16), "Level": (10, 16)}
    for sheet in book:
        count = sheet.max_column
        last = get_column_letter(count)
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
        sheet.print_title_rows = f"{HEADER_ROW}:{HEADER_ROW}"
        sheet.print_options.horizontalCentered = False
        sheet.print_area = f"A1:{last}{sheet.max_row}"
        # Titles, metadata, notes and summary are above the unchanged business table.
        write_row(sheet, 1, ("Material Requirement & Stock Report",))
        sheet.merge_cells(f"A1:{last}1")
        sheet["A1"].font = Font(name="Segoe UI", size=18, bold=True, color=INK)
        sheet.row_dimensions[1].height = 34
        for row in range(METADATA_START, METADATA_START + 5):
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
            sheet.merge_cells(start_row=row, start_column=3, end_row=row, end_column=6)
            sheet.cell(row, 1).font = Font(name="Segoe UI", size=10, bold=True, color=MUTED)
            sheet.cell(row, 1).alignment = Alignment(vertical="center", wrap_text=True)
            cell = sheet.cell(row, METADATA_VALUE_COLUMN)
            cell.font = Font(name="Segoe UI", size=10, color=INK)
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            sheet.row_dimensions[row].height = 30 if row in (5, 6) else 22
        # File names are prominent, while the original complete metadata paths remain intact.
        for row, label, source in ((3, "Storage File", stock_book.path), (5, "BOM File", bom_book.path)):
            sheet.cell(row=row, column=7, value=label).font = Font(name="Segoe UI", size=10, color=MUTED)
            # A filename beginning with '=' is also plain text, never an Excel formula.
            cell = sheet.cell(row=row + 1, column=7, value=Path(source).name)
            if cell.data_type == "f":
                cell.data_type = "s"
            sheet.merge_cells(start_row=row+1, start_column=7, end_row=row+1, end_column=count)
            cell.font = Font(name="Segoe UI", size=11, bold=True, color=INK)
            cell.alignment = Alignment(vertical="center")
        sheet.merge_cells(f"A{NOTE_ROW}:{last}{NOTE_ROW+1}")
        sheet.cell(NOTE_ROW, 1).font = Font(name="Segoe UI", size=9, color=MUTED)
        sheet.cell(NOTE_ROW, 1).alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[NOTE_ROW].height = 24
        sheet.row_dimensions[NOTE_ROW+1].height = 20
        if sheet.title == "Shortage":
            groups = ((1, "Total Materials", len(result.rows), INK),
                      (4, "Shortage", result.category_count(SHORTAGE), DARK_RED),
                      (7, "Exceptions", result.category_count(EXCEPTION), DARK_AMBER),
                      (10, "Enough", result.category_count(ENOUGH), "3B715E"))
            for col, label, value, colour in groups:
                sheet.cell(11, col, label).font = Font(name="Segoe UI", size=10, color=MUTED)
                sheet.cell(12, col, value).font = Font(name="Segoe UI", size=17, bold=True, color=colour)
                sheet.cell(12, col).alignment = Alignment(horizontal="left", vertical="center")
                sheet.merge_cells(start_row=11, start_column=col, end_row=11, end_column=col+2)
                sheet.merge_cells(start_row=12, start_column=col, end_row=12, end_column=col+2)
        else:
            write_row(sheet, 11, ("These rows require manual review.",))
            sheet.merge_cells(f"A11:{last}11")
            sheet["A11"].font = Font(name="Segoe UI", size=11, bold=True, color=DARK_AMBER)
            breakdown = Counter(code for r in rows if r.category == EXCEPTION for code in r.exceptions)
            write_row(sheet, 12, (f"Total Exceptions: {result.category_count(EXCEPTION)}",))
            sheet.merge_cells("A12:E12")
            sheet["A12"].font = Font(name="Segoe UI", size=12, bold=True, color=DARK_AMBER)
            write_row(sheet, 13, ("Types (a row may have several): " + "; ".join(f"{key}: {n}" for key, n in sorted(breakdown.items())),))
            sheet.merge_cells(f"A13:{last}13")
            sheet["A13"].font = Font(name="Segoe UI", size=9, color=MUTED)
            sheet["A13"].alignment = Alignment(wrap_text=True, vertical="center")
            sheet.row_dimensions[13].height = 30
        sheet.row_dimensions[12].height = 28
        header = [cell.value for cell in sheet[HEADER_ROW]]
        for col, label in enumerate(header, 1):
            cell = sheet.cell(HEADER_ROW, col)
            cell.font = Font(name="Segoe UI", size=10, bold=True, color=INK)
            cell.fill = PatternFill("solid", fgColor="EAF0F5")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=edge)
            low, high = limits.get(label, (12, 22))
            longest = max([len(str(label))] + [len(str(sheet.cell(r, col).value or "")) for r in range(HEADER_ROW+1, sheet.max_row+1)])
            sheet.column_dimensions[get_column_letter(col)].width = min(high, max(low, longest + 2))
        sheet.row_dimensions[HEADER_ROW].height = 32
        result_col = header.index("Result") + 1 if "Result" in header else None
        for row in range(HEADER_ROW+1, sheet.max_row+1):
            category = sheet.cell(row, result_col).value if result_col else EXCEPTION
            wrapped_lines = 1
            for col, label in enumerate(header, 1):
                cell = sheet.cell(row, col)
                cell.font = Font(name="Segoe UI", size=10, color=INK)
                cell.border = Border(bottom=edge)
                cell.alignment = Alignment(horizontal="right" if label in numeric else "left", vertical="center",
                                           wrap_text=label not in numeric and label != "Material ID")
                if cell.alignment.wrap_text and cell.value:
                    text_width = sum(2 if unicodedata.east_asian_width(char) in ('W', 'F') else 1 for char in str(cell.value))
                    column_width = sheet.column_dimensions[get_column_letter(col)].width
                    wrapped_lines = max(wrapped_lines, math.ceil(text_width / max(1, column_width * .9)))
                if category in FILLS:
                    cell.fill = FILLS[category]
                if label in numeric:
                    integral = isinstance(cell.value, int) or isinstance(cell.value, float) and cell.value.is_integer()
                    cell.number_format = "0" if integral else "General"
                if label in ("Shortage", "Calculated Shortage") and isinstance(cell.value, (int, float)) and cell.value > 0:
                    cell.font = Font(name="Segoe UI", size=10, bold=True, color=DARK_RED)
                elif label in ("Data Status", "Exception Type") and category == EXCEPTION:
                    cell.font = Font(name="Segoe UI", size=10, color=DARK_AMBER)
                elif label == "Result":
                    cell.font = Font(name="Segoe UI", size=10, bold=True, color=DARK_RED if category == SHORTAGE else DARK_AMBER if category == EXCEPTION else INK)
            sheet.row_dimensions[row].height = max(32, min(144, wrapped_lines * 14 + 8))
        sheet.freeze_panes = f"A{HEADER_ROW+1}"
        sheet.auto_filter.ref = f"A{HEADER_ROW}:{last}{max(HEADER_ROW, sheet.max_row)}"
